"""Ingestion: every format -> page-marked Markdown, index, read_section, pre-scans."""

from __future__ import annotations

import io
import json

import pytest

from bidcore.evidence import CorpusIndex
from bidcore.ingest import ingest, list_pages, read_section
from bidcore.ingest.models import PAGE_MARKER_RE


def _pdf() -> bytes:
    import pymupdf

    doc = pymupdf.open()
    texts = [
        "Request for Proposal\nSAP S/4HANA implementation for ACME Foods.\nThe programme will be delivered in three waves.",
        "Wave 1: Foods 6 months\nWave 2: Corporate + Feed\nWave 3: Retail\nThe waves run sequentially.",
        "Finance scope covers General Ledger, Asset Accounting and Controlling in Saudi Arabia.",
    ]
    for text in texts:
        page = doc.new_page()
        page.insert_text((72, 72), text, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def _docx() -> bytes:
    from docx import Document

    d = Document()
    d.add_heading("Scope of Work", 1)
    d.add_paragraph("The client runs Kronos for time management and Coupa for procurement.")
    d.add_heading("Integration Landscape", 2)
    t = d.add_table(rows=3, cols=3)
    for i, row in enumerate([["Third Party System", "Description", "SAP Module"],
                             ["Kronos", "Time and attendance", "HCM"],
                             ["Coupa", "Procurement suite", "MM"]]):
        for j, value in enumerate(row):
            t.cell(i, j).text = value
    d.add_paragraph("Data migration covers material master and customer master.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Programme Timeline"
    box = s.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.text = "Hypercare of 3 months after each go-live."
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Countries"
    ws.append(["Country", "Finance", "Sales"])
    ws.append(["Saudi Arabia", "X", "X"])
    ws.append(["UAE", "X", None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def workspace(tmp_path):
    rfp = tmp_path / "rfp"
    result = ingest([("ACME RFP.pdf", _pdf()), ("Scope.docx", _docx()), ("Timeline.pptx", _pptx()),
                     ("Matrix.xlsx", _xlsx()), ("notes.csv", b"System,Owner\nKronos,HR\n"),
                     ("readme.txt", b"Plain text requirement: SAP Analytics Cloud."),
                     ("logo.png", b"\x89PNG....")], rfp, cache_dir=tmp_path / "cache")
    return rfp, result, tmp_path


def test_every_supported_file_is_page_marked(workspace):
    rfp, result, _ = workspace
    by_name = {f.name: f for f in result.files}
    assert by_name["logo.png"].kind == "unsupported" and by_name["logo.png"].warnings
    for name in ("ACME RFP.pdf", "Scope.docx", "Timeline.pptx", "Matrix.xlsx", "notes.csv", "readme.txt"):
        record = by_name[name]
        text = (rfp.parent / record.md_path).read_text(encoding="utf-8")
        assert record.pages >= 1 and len(PAGE_MARKER_RE.findall(text)) == record.pages, name
    assert by_name["ACME RFP.pdf"].pages == 3
    assert by_name["Scope.docx"].tables == 1 and by_name["Matrix.xlsx"].tables == 1


def test_index_lists_files_outline_waves_and_third_party(workspace):
    rfp, result, tmp = workspace
    index = (rfp / "index.md").read_text(encoding="utf-8")
    assert "/rfp/acme-rfp.md" in index and "/rfp/scope.md" in index
    assert "Scope of Work" in index                       # docx heading -> outline
    assert [w["ordinal"] for w in result.declared_waves] == [1, 2, 3]
    assert result.sequencing_hint == "sequential"
    assert any(a["page"] == 2 for a in result.wave_anchors)
    names = {c["name"] for c in result.third_party_candidates}
    assert {"Kronos", "Coupa"} <= names
    assert json.loads((tmp / "cache" / "prescan.json").read_text())["declared_waves"]


def test_read_section_pages_and_heading(workspace):
    rfp, _, _ = workspace
    assert list_pages(rfp, "acme-rfp.md") == [1, 2, 3]
    two = read_section(rfp, "/rfp/acme-rfp.md", 2)
    assert two.startswith("<!-- page: 2 -->") and "Wave 2" in two and "Finance scope" not in two
    assert "Finance scope" in read_section(rfp, "acme-rfp", 2, 3)
    block = read_section(rfp, "scope.md", heading="integration landscape")
    assert "Kronos" in block and "Scope of Work" not in block
    assert "No pages" in read_section(rfp, "acme-rfp.md", 9)
    with pytest.raises(FileNotFoundError):
        read_section(rfp, "index.md", 1)


def test_reingest_adds_file_and_keeps_index(workspace):
    rfp, _, _ = workspace
    result = ingest([("Addendum.txt", b"Addendum: Wave 4: Logistics")], rfp)
    assert [f.slug for f in result.files] == ["addendum"]
    index = (rfp / "index.md").read_text(encoding="utf-8")
    assert "/rfp/acme-rfp.md" in index and "/rfp/addendum.md" in index
    data = _pdf()                                        # PDF bytes differ per build (ids/dates)
    first = ingest([("ACME RFP.pdf", data)], rfp).files[0].slug
    assert first == "acme-rfp-2"                         # different bytes, same name -> new slug
    assert ingest([("ACME RFP.pdf", data)], rfp).files[0].slug == first    # same bytes -> reused


def test_evidence_grounding_reads_the_workspace(workspace):
    rfp, _, _ = workspace
    corpus = CorpusIndex.from_dir(rfp)
    assert corpus.contains("General Ledger, Asset Accounting and Controlling")
    assert corpus.contains("Kronos for time management")
    assert not corpus.contains("Ariba sourcing")


def test_captioner_is_called_once_per_image_and_cached(tmp_path):
    from PIL import Image, ImageDraw
    from docx import Document
    from docx.shared import Inches

    img = Image.new("RGB", (600, 400), "white")
    draw = ImageDraw.Draw(img)
    for i in range(0, 600, 7):
        draw.line([(i, 0), (600 - i, 400)], fill=(i % 255, 80, 160))
    png = io.BytesIO()
    img.save(png, format="PNG")
    d = Document()
    d.add_paragraph("Integration landscape below")
    d.add_picture(io.BytesIO(png.getvalue()), width=Inches(4))
    buf = io.BytesIO()
    d.save(buf)
    calls = []

    def caption(data, mime, context):
        calls.append(context)
        return "Image type: diagram\nKronos <-> SAP HCM : time data"

    first = ingest([("land.docx", buf.getvalue())], tmp_path / "rfp", caption=caption, cache_dir=tmp_path / "c")
    assert first.files[0].images_captioned == 1 and len(calls) == 1
    text = (tmp_path / "rfp" / "land.md").read_text(encoding="utf-8")
    assert "[EMBEDDED IMAGE -- p.1]" in text and "Kronos <-> SAP HCM" in text
    ingest([("land.docx", buf.getvalue())], tmp_path / "rfp2", caption=caption, cache_dir=tmp_path / "c")
    assert len(calls) == 1                               # cache hit

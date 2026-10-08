"""Effort workbook: render -> formulas evaluate to the sizing figures; reviewer edits round-trip (call 2)."""

from __future__ import annotations

import shutil
import subprocess

import pytest
from openpyxl import load_workbook

from bidcore.render.workbook import WorkbookError, apply_overrides, read_bid_sheet, read_edits, render_workbook
from bidcore.sizing import size_bid
from tests.fixtures_bid import build_ledger


@pytest.fixture()
def rendered(tmp_path):
    ledger = build_ledger()
    sizing = size_bid(ledger)
    path = render_workbook(ledger, sizing, tmp_path / "out" / "effort.xlsx", manifest_dir=tmp_path / "ledger")
    return ledger, sizing, path, tmp_path


def _find(ws, label, col=1):
    for row in range(1, ws.max_row + 1):
        if str(ws.cell(row=row, column=col).value or "").strip() == label:
            return row
    raise AssertionError(f"{label!r} not found in {ws.title}")


def _recalculated(path, tmp_path):
    """Evaluate every formula with LibreOffice (headless) and return a data_only workbook."""
    out = tmp_path / "recalc"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(path)],
                   check=True, capture_output=True, timeout=180)
    return load_workbook(out / path.name, data_only=True)


def test_sheets_order_and_bid_stamp(rendered):
    ledger, _, path, _ = rendered
    wb = load_workbook(path)
    names = wb.sheetnames
    assert names[:3] == ["Summary of Project Effort", "Functional Scope Estimation", "Project Timeline"]
    assert {"Finance", "Sourcing and Procurement", "R&D_Engineering"} <= set(names)
    assert names[-1] == "_bid" and wb["_bid"].sheet_state == "hidden"
    stamp = read_bid_sheet(path)
    assert stamp["bid_id"] == ledger.meta.bid_id and stamp["render_id"]


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_formulas_evaluate_to_sizing_figures(rendered):
    _, sizing, path, tmp = rendered
    wb = _recalculated(path, tmp)
    ws = wb["Summary of Project Effort"]
    s = sizing.summary
    expected = {"Functional Scope": s.functional_scope, "Data Migration Efforts": s.data_migration,
                "Tech Dev Scope": s.tech_dev, "Security": s.security, "BASIS & SolMan Efforts": s.basis,
                "Analytics Efforts": s.analytics,
                "Total Build Effort(Upto Realization Phase)": s.total_build, "Final preparation phase": s.final_prep,
                "Go-Live Effort": s.go_live, "Project Mangement Effort": s.project_mgmt, "Hypercare": s.hypercare,
                "Summary of Imp": s.summary_of_imp, "Risk Contingency": s.risk_contingency,
                "Total Project Effort": s.total_project_effort, "Total Project Cost (USD)": s.total_project_cost_usd}
    for label, value in expected.items():
        assert ws.cell(row=_find(ws, label), column=2).value == pytest.approx(value, abs=0.011), label
    timeline = wb["Project Timeline"]
    grid_pd = timeline.cell(row=_find(timeline, "Grid rows sum (all waves x days-per-month)"), column=2).value
    assert grid_pd == pytest.approx(sizing.plan.grand_total_mm * 21, abs=0.5)
    total = timeline.cell(row=_find(timeline, "Sum: Total Project Effort"), column=2).value
    assert total == pytest.approx(s.total_project_effort, abs=0.011)


def test_unedited_workbook_has_no_edits(rendered):
    _, _, path, tmp = rendered
    manifest, edits = read_edits(path, tmp / "ledger")
    assert edits == [] and manifest.tables


def test_reviewer_edits_round_trip(rendered):
    ledger, sizing, path, tmp = rendered
    wb = load_workbook(path)
    fin = wb["Finance"]
    first = 2                                                  # module sheets: header row 1 (old layout)
    line_id = fin.cell(row=first, column=15).value
    fin.cell(row=first, column=12, value=2)                    # multiplication factor
    new_wave = "Wave 1" if fin.cell(row=first, column=14).value == "Wave 2" else "Wave 2"
    fin.cell(row=first, column=14, value=new_wave)             # wave tag
    fin.cell(row=first + 1, column=7, value=99)                # workshops & configuration
    deleted_line = fin.cell(row=first + 2, column=15).value
    fin.delete_rows(first + 2)                                 # delete a catalogue line
    nc = wb["Non Catalogue SAP Tools"]
    nc.cell(row=3, column=4, value=40)                         # non-catalogue effort (row 2 = block header)
    nc.insert_rows(4)
    nc.cell(row=4, column=1, value="SAP Signavio")
    nc.cell(row=4, column=4, value=12)                         # added row
    basis = wb["Basis Scope"]
    basis.cell(row=3, column=3, value=1)                       # DEV effort (sync may have raised it)
    sec = wb["Security Scope"]
    sec.delete_rows(4)                                         # drop GRC Access Control
    td = wb["Tech Dev Scope"]
    td.cell(row=2, column=13, value=123)                       # overwrite a formula -> conflict
    summary = wb["Summary of Project Effort"]
    summary.cell(row=_find(summary, "Daily Rate (USD)", 4), column=5, value=400)
    path2 = tmp / "reviewed.xlsx"
    wb.save(path2)

    _, edits = read_edits(path2, tmp / "ledger")
    kinds = {(e.section, e.kind, e.field) for e in edits}
    assert ("scope_items", "edit", "multiplication_factor") in kinds
    assert ("scope_items", "edit", "wave") in kinds
    assert ("scope_items", "deleted_row", "") in kinds
    assert ("non_catalogue", "added_row", "") in kinds
    assert ("security", "deleted_row", "") in kinds
    assert ("tech_dev", "conflict", "total") in kinds
    assert ("", "setting", "daily_rate_usd") in kinds

    summary_text = apply_overrides(ledger, edits)
    assert "conflict" in summary_text
    assert [r.row_id for r in ledger.security.rows] == ["sec-1"]
    assert ledger.basis.rows[0].dev == 1
    assert {r.name for r in ledger.non_catalogue.rows} == {"SAP Solution Manager 7.2", "SAP Signavio"}
    assert ledger.non_catalogue.rows[0].effort_days == 40

    after = size_bid(ledger)
    lines = {r.row_id: r for m in after.effort.modules.values() for r in m.rows}
    assert deleted_line not in lines
    assert lines[line_id].multiplication_factor == 2 and lines[line_id].wave == new_wave
    assert after.summary.daily_rate_usd == 400
    assert after.summary.security == pytest.approx(20, abs=30)       # scope sync may move it within its band
    assert after.effort.non_catalogue_effort == 52


def test_foreign_workbook_is_rejected(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.save(tmp_path / "old.xlsx")
    with pytest.raises(WorkbookError):
        read_edits(tmp_path / "old.xlsx", tmp_path)

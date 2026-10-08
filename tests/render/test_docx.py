"""Proposal .docx: cover tokens, outline order, placeholder resolution, disclosure, missing drafts."""

from __future__ import annotations

from docx import Document

from bidcore.figures import compute_figures
from bidcore.ledger.sections_proposal import ClientRequirement, Disclosure, ResponseRequirements
from bidcore.outline import merged_outline
from bidcore.render.docx import render_proposal
from bidcore.render.docx.builder import MISSING_DRAFT_TEXT, WITHHELD_FIGURE_TEXT
from bidcore.sizing import size_bid
from tests.fixtures_bid import build_ledger


def _text(doc) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


def _all_text(doc) -> str:
    return "".join(t.text or "" for t in doc.element.iter() if t.tag.endswith("}t"))


def _render(tmp_path, disclosure: Disclosure | None = None):
    ledger = build_ledger()
    ledger.response_requirements.data = ResponseRequirements(
        requirements=[ClientRequirement(title="Indicative effort by wave", intent="Show effort per wave",
                                        kind="indicative_breakdown", group_by="wave",
                                        placement_after_section_id="6.1")],
        disclosure=disclosure or Disclosure.full())
    sizing = size_bid(ledger)
    ledger.figures = compute_figures(ledger, sizing)
    drafts = tmp_path / "drafts"
    drafts.mkdir()
    (drafts / "1.1.md").write_text(
        "ACME wants **one global template**.\n\n- Finance first\n- then procurement\n\n"
        "The programme runs {{fig:programme_months}} with a total of {{fig:total_project_effort}}.\n", encoding="utf-8")
    (drafts / "4.3.md").write_text("Our plan:\n\n{{table:wave_plan}}\n\nThe diagram follows.", encoding="utf-8")
    (drafts / "6.2.md").write_text("The price is {{fig:total_project_cost}}.\n\n| A | B |\n|---|---|\n| x | y |",
                                   encoding="utf-8")
    out = render_proposal(ledger, sizing, drafts, tmp_path / "out.docx")
    return ledger, sizing, Document(str(out))


def test_cover_tokens_and_toc_refresh(tmp_path):
    _, _, doc = _render(tmp_path)
    text = _all_text(doc)
    assert "{{CLIENT_NAME}}" not in text and "{{DATE}}" not in text and "ACME Foods" in text
    assert doc.settings.element.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}updateFields") is not None


def test_sections_in_outline_order_with_client_section(tmp_path):
    ledger, _, doc = _render(tmp_path)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    titles = [s.title for s in merged_outline(ledger)]
    titles = [f"{s.id} {s.title}" if s.client_required else s.title for s in merged_outline(ledger)]
    assert [h for h in headings if h in titles] == titles
    assert headings.index("6.1.1 Indicative effort by wave") == headings.index("Effort Estimation") + 1


def test_placeholders_resolved_and_artifacts_placed(tmp_path):
    _, sizing, doc = _render(tmp_path)
    text = _text(doc)
    assert "{{" not in text
    assert f"{sizing.summary.total_project_effort:,.2f}".rstrip("0").rstrip(".") in text
    assert MISSING_DRAFT_TEXT in text                       # sections without drafts are flagged
    table_titles = [p.text for p in doc.paragraphs if p.style.name == "Caption"]
    assert table_titles.count("Delivery waves") == 1          # placed inline, not repeated as the artifact
    assert "Effort estimate" in table_titles and "Cost estimate (USD)" in table_titles
    assert len(doc.inline_shapes) == 3                         # methodology, architecture, timeline


def test_disclosure_withholds_commercials(tmp_path):
    _, _, doc = _render(tmp_path, Disclosure.full().model_copy(update={"commercial_detail": False, "effort_detail": False}))
    text = _text(doc)
    assert WITHHELD_FIGURE_TEXT in text
    captions = [p.text for p in doc.paragraphs if p.style.name == "Caption"]
    assert "Cost estimate (USD)" not in captions and "Effort estimate" not in captions
    assert "Indicative effort by wave" not in captions

"""Call 2 end to end with scripted models: reviewed workbook -> edits applied -> proposal team drafts ->
coverage retries -> .docx with figures from the edited ledger."""

from __future__ import annotations

import asyncio

from docx import Document
from openpyxl import load_workbook

from bidcore.figures import compute_figures
from bidcore.render.workbook import render_workbook
from bidcore.sizing import size_bid
from harness import llm
from harness.context import RunContext
from harness.workspace import BidWorkspace
from tests.fixtures_bid import build_ledger
from tests.harness.fakes import call, factory, turn
from tests.harness.test_effort_e2e import RFP, harness_env  # noqa: F401  (fixture)


def _call1_outputs(tmp):
    ws = BidWorkspace.open("fixture01", base=tmp / "ws")
    ledger = build_ledger()
    sizing = size_bid(ledger)
    ledger.figures = compute_figures(ledger, sizing)
    ws.ledger.save(ledger)
    book = render_workbook(ledger, sizing, ws.outputs / "effort.xlsx", manifest_dir=ws.ledger.dir)
    wb = load_workbook(book)
    summary = wb["Summary of Project Effort"]
    for row in range(1, summary.max_row + 1):
        if summary.cell(row=row, column=4).value == "Daily Rate (USD)":
            summary.cell(row=row, column=5, value=350)
    reviewed = tmp / "reviewed.xlsx"
    wb.save(reviewed)
    rfp = tmp / "acme.txt"
    rfp.write_text(RFP, encoding="utf-8")
    return ws, reviewed, rfp


def proposal_scripts() -> dict:
    return {
        "proposal-orchestrator": [
            turn(call("outline")),
            turn(call("task", subagent_type="requirements-analyst", description="Record response requirements.")),
            turn(call("task", subagent_type="section-writer", description="Draft sections 1.1 and 6.2.")),
            turn(call("drafts_status")),
            turn(text="Drafted what I could."),
        ],
        "requirements-analyst": [
            turn(call("ledger_write_response_requirements", data={
                "requirements": [{"title": "Indicative effort by wave", "intent": "Effort per wave",
                                  "kind": "indicative_breakdown", "group_by": "wave",
                                  "placement_after_section_id": "6.1"}],
                "disclosure": {"resource_location": False, "source": "rfp"}})),
            turn(text="1 requirement."),
        ],
        "section-writer": [
            turn(call("bid_facts", view="figures")),
            turn(call("write_draft", section_id="1.1",
                      markdown="ACME Foods wants one **global template** for finance."),
                 call("write_draft", section_id="6.2", markdown="Our price is {{fig:total_project_cost}}."),
                 call("write_draft", section_id="6.1.1", markdown="The table below shows the split by wave.")),
            turn(text="3 drafts."),
        ],
    }


def test_proposal_call_end_to_end(harness_env):  # noqa: F811
    tmp, log = harness_env
    from workflows.proposal import build_proposal_workflow

    ws, reviewed, rfp = _call1_outputs(tmp)
    llm.set_model_factory(factory(proposal_scripts(), log))
    run = RunContext(ws=ws, team="proposal")
    flow = build_proposal_workflow(run)
    state = asyncio.run(flow.ainvoke({"workbook": str(reviewed), "rfp_uploads": [str(rfp)],
                                      "instructions": "Stress local presence."},
                                     config={"configurable": {"thread_id": "t"}, "recursion_limit": 100}))

    ledger = ws.ledger.load()
    assert any(o.kind == "setting" and o.field == "daily_rate_usd" and o.applied for o in ledger.overrides)
    assert (ws.notes / "reviewer-edits.md").is_file() and (ws.notes / "presales-instructions.md").is_file()
    assert ledger.response_requirements.data.disclosure.resource_location is False
    assert {"1.1", "6.2", "6.1.1"} <= {p.stem for p in ws.drafts.glob("*.md")}
    assert state["missing"]                                  # the scripted writer skipped most sections
    doc = Document(state["document"])
    text = "\n".join(p.text for p in doc.paragraphs)
    cost = ledger.figures["fig"]["total_project_cost"]
    assert f"Our price is {cost}." in text
    assert size_bid(ledger).summary.daily_rate_usd == 350    # the reviewer's rate reached the document
    captions = [p.text for p in doc.paragraphs if p.style.name == "Caption"]
    assert "Indicative effort by wave" in captions
    headers = [[c.text for c in t.rows[0].cells] for t in doc.tables]
    assert ["Role", "Waves"] in headers                      # roster kept, Location column withheld
    assert not any("Location" in h for h in headers)

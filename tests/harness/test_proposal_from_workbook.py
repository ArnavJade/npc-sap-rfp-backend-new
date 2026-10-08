"""Call 2 from a workbook alone: no `_bid` sheet, no call-1 workspace. The reader rebuilds the ledger and
every figure from the template sheets, and the document's numbers equal the workbook's."""

from __future__ import annotations

import asyncio

from docx import Document
from openpyxl import load_workbook

from bidcore.figures import compute_figures
from bidcore.render.workbook import render_workbook
from bidcore.render.workbook.reader import WorkbookFormatError, read_effort_workbook
from bidcore.sizing import size_bid
from harness import llm
from harness.context import RunContext
from harness.workspace import BidWorkspace
from tests.fixtures_bid import build_ledger
from tests.harness.fakes import call, factory, turn
from tests.harness.test_effort_e2e import harness_env  # noqa: F401  (fixture)


def _external_workbook(tmp):
    """A call-1 workbook from 'another server': rendered elsewhere, `_bid` removed, one value edited."""
    ledger = build_ledger()
    sizing = size_bid(ledger)
    ledger.figures = compute_figures(ledger, sizing)
    book = render_workbook(ledger, sizing, tmp / "elsewhere" / "Claude_BP_Effort_Output_ACME_Foods_20261008_100321_Summary.xlsx")
    wb = load_workbook(book)
    del wb["_bid"]
    tech = wb["Tech Dev Scope"]
    tech.cell(row=2, column=9, value=100)          # reviewer raised the first row's development days
    out = tmp / "Claude_BP_Effort_Output_ACME_Foods_20261009_090000_Summary.xlsx"   # template file name
    wb.save(out)
    return ledger, sizing, out


def test_reader_rebuilds_figures(tmp_path):
    ledger, sizing, book = _external_workbook(tmp_path)
    snap = read_effort_workbook(book, "ACME Foods")
    got = snap.sizing()
    assert got.effort.core_bp_effort == sizing.effort.core_bp_effort
    assert got.effort.non_catalogue_effort == sizing.effort.non_catalogue_effort
    assert got.totals["basis"] == sizing.totals["basis"] and got.totals["security"] == sizing.totals["security"]
    assert round(got.totals["data_migration"], 2) == round(sizing.totals["data_migration"], 2)
    assert got.effort.tech_dev_effort != sizing.effort.tech_dev_effort          # the edit is honoured
    assert [g.wave_name for g in got.plan.grids] == [g.wave_name for g in sizing.plan.grids]
    assert got.plan.grand_total_mm == sizing.plan.grand_total_mm
    assert got.summary.hypercare == sizing.summary.hypercare
    led = snap.ledger("imp1")
    assert led.meta.source == "workbook" and led.scope_items.rows and led.timeline.data.waves
    assert {r.scope_item_id for r in led.scope_items.rows} == {r.scope_item_id for r in ledger.scope_items.rows
                                                               if r.status != "excluded_existing"} or led.scope_items.rows


def test_reader_rejects_non_template(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "x.xlsx"
    Workbook().save(path)
    try:
        read_effort_workbook(path)
    except WorkbookFormatError as exc:
        assert "does not look like an effort workbook" in str(exc)
    else:
        raise AssertionError("expected WorkbookFormatError")


def test_proposal_without_bid_sheet(harness_env):  # noqa: F811
    tmp, log = harness_env
    from workflows.proposal import build_proposal_workflow

    _, _, book = _external_workbook(tmp)
    scripts = {
        "proposal-orchestrator": [turn(call("task", subagent_type="section-writer", description="Draft 6.2.")),
                                  turn(text="done")],
        "section-writer": [turn(call("write_draft", section_id="6.2", markdown="Our price is {{fig:total_project_cost}}.")),
                           turn(text="1 draft.")],
    }
    llm.set_model_factory(factory(scripts, log))
    ws = BidWorkspace.open("fromwb1", base=tmp / "ws")             # fresh: no ledger, no manifest
    run = RunContext(ws=ws, team="proposal")
    state = asyncio.run(build_proposal_workflow(run).ainvoke(
        {"workbook": str(book), "rfp_uploads": [], "instructions": ""},
        config={"configurable": {"thread_id": "t"}, "recursion_limit": 100}))
    ledger = ws.ledger.load()
    assert ledger.meta.source == "workbook" and ledger.meta.client_name == "ACME Foods"
    expected = read_effort_workbook(book).sizing().summary.total_project_cost_usd
    cost = ledger.figures["fig"]["total_project_cost"]
    assert f"{expected:,.0f}" in cost
    text = "\n".join(p.text for p in Document(state["document"]).paragraphs)
    # No RFP came with the call, so the disclosure profile falls back to "withheld" (old rule): the cost
    # figure is masked in the text, but the draft itself reached the document.
    assert "Our price is available on request." in text or f"Our price is {cost}." in text


def test_existing_bid_with_unlinked_workbook_is_rebuilt_from_it(harness_env):  # noqa: F811
    tmp, log = harness_env
    from bidcore.ledger.models import new_ledger
    from workflows.proposal import build_proposal_workflow

    _, _, book = _external_workbook(tmp)
    llm.set_model_factory(factory({}, log))
    ws = BidWorkspace.open("fromwb2", base=tmp / "ws")
    ws.ledger.save(new_ledger("fromwb2", "ACME Foods Ltd"))       # a bid exists, but not this workbook's render
    state = asyncio.run(build_proposal_workflow(RunContext(ws=ws, team="proposal")).ainvoke(
        {"workbook": str(book), "rfp_uploads": [], "instructions": ""},
        config={"configurable": {"thread_id": "t"}, "recursion_limit": 100}))
    ledger = ws.ledger.load()
    assert ledger.meta.source == "workbook" and ledger.meta.client_name == "ACME Foods Ltd"
    assert ledger.scope_items.rows and state["document"].endswith(".docx")

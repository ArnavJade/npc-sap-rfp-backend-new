"""Workbook styling (colours and fonts copied from the old writers / effort-estimation template) and
small helpers to clone a prototype row's style onto generated rows."""

from __future__ import annotations

from copy import copy

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet


def fill(hex_colour: str) -> PatternFill:
    return PatternFill(start_color=hex_colour, end_color=hex_colour, fill_type="solid")


THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = fill("1F4E78")                  # module / summary headers
HEADER_FONT = Font(color="FFFFFF", bold=True)
NAVY_FILL = fill("002060")                    # scope-sheet headers, Summary of Project Effort labels
TITLE_FILL = fill("D0CECE")                   # scope-sheet title bands
AMBER_FILL = fill("FFC000")                   # Summary title, 'Master Data' band
GREEN_FILL = fill("92D050")                   # Summary total rows
WAVE_FILL = fill("DDEBF7")                    # Delivery-wave rows
GRID_FILL = fill("FCE4D6")                    # resource-plan block titles
TOOLS_FILL = fill("FFF2CC")                   # non-catalogue block header
BOLD = Font(bold=True)
ESTIMATE_FONT = Font(italic=True, color="806000")
ESTIMATE_FORMAT = '"~"General" (est.)"'
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(wrap_text=True, vertical="center")


def row_styles(ws: Worksheet, row: int, max_col: int) -> list:
    """Snapshot one prototype row's cell styles (StyleArray copies)."""
    return [copy(ws.cell(row=row, column=c)._style) for c in range(1, max_col + 1)]


def apply_row_styles(ws: Worksheet, row: int, styles: list) -> None:
    for col, style in enumerate(styles, start=1):
        ws.cell(row=row, column=col)._style = copy(style)


def clear_rows(ws: Worksheet, first: int, last: int, max_col: int) -> None:
    """Remove prototype rows' content, styles and merges before generated rows are written."""
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row >= first and merged.max_row <= last:
            ws.unmerge_cells(str(merged))
    for r in range(first, last + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.value = None
            cell._style = copy(ws.cell(row=last + 50, column=c)._style)   # a never-styled cell
        if r in ws.row_dimensions:
            ws.row_dimensions[r].height = None


def header_row(ws: Worksheet, row: int, headers: list[str], fill_: PatternFill = HEADER_FILL,
               font: Font = HEADER_FONT, bordered: bool = False, centered: bool = False) -> None:
    for c, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=text)
        cell.fill, cell.font = fill_, font
        cell.alignment = CENTER if centered else Alignment(wrap_text=True, vertical="center")
        if bordered:
            cell.border = BORDER

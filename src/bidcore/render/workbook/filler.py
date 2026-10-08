"""Write the effort workbook from the ledger + sizing result (no agent ever edits the file).

Layouts are ported from the old writers - effort_calculator_layer4.write_effort_workbook (Summary of
Project Effort, Functional Scope Estimation, Project Timeline, module sheets, Non Catalogue SAP
Tools, Tech Dev Scope) and data_migration_scope / basis_scope / security_scope / analytics_scope -
with the same titles, headers, fills, row positions and live formulas. Tab order (old
_reorder_sheets_for_client): Summary of Project Effort, Functional Scope Estimation, Project Timeline,
one sheet per LOB, Non Catalogue SAP Tools, Tech Dev Scope, Data Migration Scope, Basis Scope,
Security Scope, Analytics Scope, then the hidden `_bid` sheet.

Additions over the old workbook: a hidden row-id column after every editable table, the hidden `_bid`
sheet, and the daily / hypercare rates plus a Total Project Cost row on the Summary of Project Effort.
Every table row is recorded in the render manifest, which is how call 2 reads the edits back.
The old hidden 'Pipeline Data' sheet is not written: call 2 reads the ledger, not the workbook.
"""

from __future__ import annotations

import uuid
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from bidcore.ledger.models import Ledger
from bidcore.policy import Policy, get_policy
from bidcore.render.workbook.layout import (
    FORMULA, GRID_FIRST_MONTH_COL, GRID_SKIP_LABELS, GRID_TOTAL_LABEL, HYPERCARE_LABEL, HYPERCARE_SETTING, INPUT,
    SUMMARY_PARAMETERS, TableLayout, module_sheet_name, quote_sheet, table_layouts, xl_criteria,
)
from bidcore.render.workbook.manifest import BID_SHEET, ColumnRef, RenderManifest, TableManifest, save_manifest
from bidcore.sizing import SizingResult


def _fill(hex_colour: str) -> PatternFill:
    return PatternFill(start_color=hex_colour, end_color=hex_colour, fill_type="solid")


HEADER_FILL, HEADER_FONT, TOTAL_FONT = _fill("1F4E78"), Font(color="FFFFFF", bold=True), Font(bold=True)
NAVY, TITLE_GREY, AMBER, GREEN = _fill("002060"), _fill("D0CECE"), _fill("FFC000"), _fill("92D050")
WAVE_FILL, GRID_FILL, TOOLS_FILL = _fill("DDEBF7"), _fill("FCE4D6"), _fill("FFF2CC")
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(wrap_text=True, vertical="center")
ESTIMATE_FONT = Font(italic=True, color="806000")
ESTIMATE_FORMAT = '"~"General" (est.)"'
DM_WAVE_HEADERS = ("Wave", "Sequence", "Partition Axis", "Duration (weeks)", "Hypercare (weeks)", "Countries")
IN_SCOPE, OUT_OF_SCOPE = "In Scope", "Out of Scope"


def _round(value: Any) -> Any:
    return round(float(value), 2) if isinstance(value, (int, float)) and not isinstance(value, bool) else value


def _autosize(ws: Worksheet) -> None:
    """Old _autosize: width = longest value + 2, between 10 and 45."""
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 2, 10), 45)


def _header(ws: Worksheet, headers: list[str], row: int = 1) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, WRAP


def _titled_header(ws: Worksheet, title: str, layout: TableLayout) -> None:
    """Scope-sheet top: merged grey title band (row 1), navy header row (row 2), widths."""
    last = len(layout.columns)
    for col, column in enumerate(layout.columns, start=1):
        if column.width:
            ws.column_dimensions[get_column_letter(col)].width = column.width
        ws.cell(row=1, column=col).border = BORDER
    cell = ws.cell(row=1, column=1, value=title)
    cell.fill, cell.font, cell.alignment = TITLE_GREY, TOTAL_FONT, CENTER
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
    for col, column in enumerate(layout.columns, start=1):
        cell = ws.cell(row=2, column=col, value=column.header)
        cell.fill, cell.font, cell.alignment, cell.border = NAVY, HEADER_FONT, CENTER, BORDER
    ws.row_dimensions[2].height = 29


def _hide_row_ids(ws: Worksheet, layout_or_col: TableLayout | int) -> None:
    col = layout_or_col if isinstance(layout_or_col, int) else layout_or_col.row_id_col
    letter = get_column_letter(col)
    ws.column_dimensions[letter].hidden = True
    ws.column_dimensions[letter].width = 18


class _Table:
    """Writes one table's data rows (with hidden row ids) and records them in the manifest."""

    def __init__(self, ws: Worksheet, layout: TableLayout, manifest: RenderManifest, header_row: int,
                 first_row: int, wave: str = "", bordered: bool = False):
        self.ws, self.layout, self.row, self.first, self.bordered = ws, layout, first_row, first_row, bordered
        self.entry = TableManifest(
            key=layout.key, section=layout.section, sheet=ws.title, header_row=header_row,
            label_col=layout.label_col, total_label=layout.total_label, row_id_col=layout.row_id_col,
            columns={i: ColumnRef(field=c.field, kind=c.kind) for i, c in enumerate(layout.columns, start=1)},
            wave=wave, skip_labels=list(layout.skip_labels))
        manifest.tables.append(self.entry)

    def add(self, row_id: str, values: dict[str, Any], formats: dict[int, tuple] | None = None) -> int:
        r = self.row
        recorded: dict[str, Any] = {}
        for c, column in enumerate(self.layout.columns, start=1):
            if c in self.layout.row_formulas:
                cell = self.ws.cell(row=r, column=c, value=self.layout.row_formulas[c].format(r=r))
            else:
                value = _round(values.get(column.field))
                cell = self.ws.cell(row=r, column=c, value=value)
                if column.kind != FORMULA:
                    recorded[column.field] = value
            if self.bordered:
                cell.border = BORDER
                cell.alignment = CENTER if c > 1 else WRAP
            if formats and c in formats:
                cell.number_format, cell.font = formats[c]
        self.ws.cell(row=r, column=self.layout.row_id_col, value=row_id)
        self.entry.rows[row_id] = recorded
        self.row += 1
        return r

    def total_row(self, label: str | None = None, bordered: bool = False) -> int:
        """The Total row: label in the layout's label column, SUM over each total column."""
        r = self.row
        self.ws.cell(row=r, column=self.layout.label_col, value=label or self.layout.total_label).font = TOTAL_FONT
        last = r - 1
        for col in self.layout.total_cols:
            letter = get_column_letter(col)
            value = f"=SUM({letter}{self.first}:{letter}{last})" if last >= self.first else 0
            cell = self.ws.cell(row=r, column=col, value=value)
            cell.font = TOTAL_FONT
            if bordered:
                cell.alignment = CENTER
        if bordered:
            for col in range(1, len(self.layout.columns) + 1):
                self.ws.cell(row=r, column=col).border = BORDER
                self.ws.cell(row=r, column=col).font = TOTAL_FONT
        self.row += 1
        return r


# --------------------------------------------------------------------------- module (LOB) sheets

def _module_sheets(wb: Workbook, layout: TableLayout, sizing: SizingResult, manifest: RenderManifest,
                   used: set[str]) -> dict[str, tuple[str, int, int]]:
    """One sheet per LOB with lines; returns LOB -> (sheet title, last data row, TOTAL row)."""
    out: dict[str, tuple[str, int, int]] = {}
    for lob, module in sizing.effort.modules.items():
        if not module.rows:
            continue
        ws = wb.create_sheet(module_sheet_name(lob, used))
        _header(ws, [c.header for c in layout.columns])
        table = _Table(ws, layout, manifest, header_row=1, first_row=2)
        for row in module.rows:
            table.add(row.row_id, {**row.model_dump(), "description": row.description})
        last = table.row - 1
        total = table.total_row()
        _autosize(ws)
        _hide_row_ids(ws, layout)
        manifest.lob_sheets[ws.title] = lob
        out[lob] = (ws.title, last, total)
    return out


def _functional_scope(ws: Worksheet, headers: list[str], modules: dict[str, tuple[str, int, int]],
                      sizing: SizingResult, labels: dict[str, str]) -> int | None:
    """Old Functional Scope Estimation: per LOB a bold header row (BP ID count, link to the module
    sheet TOTAL) and Business Area sub-rows (COUNTIF / SUMIF over the module sheet, by effort desc);
    then 'Total Functional Scope' = the LOB header rows. Returns that total row."""
    _header(ws, headers)
    r = 2
    header_rows: list[int] = []
    for lob in sorted(modules):
        title, last, total = modules[lob]
        ref = quote_sheet(title)
        ws.cell(row=r, column=1, value=labels.get(lob, lob)).font = TOTAL_FONT
        ws.cell(row=r, column=4, value=f"={ref}!M{total}").font = TOTAL_FONT
        header_rows.append(r)
        head = r
        r += 1
        effort: dict[str, float] = defaultdict(float)
        for line in sizing.effort.modules[lob].rows:
            effort[line.business_area or "Other"] += line.total
        ba_col, eff_col = f"{ref}!$B$2:$B${last}", f"{ref}!$M$2:$M${last}"
        first = r
        for area in sorted(effort, key=lambda a: -effort[a]):
            crit = [xl_criteria(area)] + ([xl_criteria("")] if area == "Other" else [])
            ws.cell(row=r, column=2, value=area)
            ws.cell(row=r, column=3, value="=" + "+".join(f"COUNTIF({ba_col},{c})" for c in crit))
            ws.cell(row=r, column=4, value="=" + "+".join(f"SUMIF({ba_col},{c},{eff_col})" for c in crit))
            r += 1
        cell = ws.cell(row=head, column=3, value=f"=SUM(C{first}:C{r - 1})" if r > first else 0)
        cell.font = TOTAL_FONT
    total = None
    if header_rows:
        ws.cell(row=r, column=1, value="Total Functional Scope").font = TOTAL_FONT
        ws.cell(row=r, column=4, value="=" + "+".join(f"D{x}" for x in header_rows)).font = TOTAL_FONT
        total = r
    _autosize(ws)
    return total


def _non_catalogue(wb: Workbook, title: str, layout: TableLayout, sizing: SizingResult,
                   manifest: RenderManifest) -> int | None:
    """Old Non Catalogue SAP Tools sheet (only when there are tools); returns its Total row."""
    if not sizing.effort.non_catalogue:
        return None
    ws = wb.create_sheet(title)
    _header(ws, [c.header for c in layout.columns])
    ws.cell(row=2, column=1, value="Non-Catalogue SAP Tools").font = TOTAL_FONT
    for col, value in ((1, None), (2, ""), (3, "-"), (4, "-")):
        cell = ws.cell(row=2, column=col)
        if value is not None:
            cell.value = value
        cell.fill = TOOLS_FILL
    table = _Table(ws, layout, manifest, header_row=1, first_row=3)
    for item in sizing.effort.non_catalogue:
        table.add(item.row_id, {"name": item.name, "description": item.description or item.label or "",
                                "bp_id_count": "-", "effort_days": item.effort_days},
                  formats={4: (ESTIMATE_FORMAT, ESTIMATE_FONT)})
    table.row += 1                       # one blank row before the total, as the old sheet
    first, last = table.first, table.row - 2
    r = table.row
    ws.cell(row=r, column=1, value="Total Non-Catalogue SAP Tools").font = TOTAL_FONT
    ws.cell(row=r, column=4, value=f"=SUM(D{first}:D{last})").font = TOTAL_FONT
    _autosize(ws)
    _hide_row_ids(ws, layout)
    return r


def _tech_dev(wb: Workbook, title: str, layout: TableLayout, sizing: SizingResult,
              manifest: RenderManifest) -> int | None:
    ws = wb.create_sheet(title)
    _header(ws, [c.header for c in layout.columns])
    table = _Table(ws, layout, manifest, header_row=1, first_row=2)
    for row in sizing.effort.tech_dev:
        table.add(row.row_id, row.model_dump())
        ws.cell(row=table.row - 1, column=13).font = TOTAL_FONT
    total = table.total_row() if sizing.effort.tech_dev else None
    _autosize(ws)
    _hide_row_ids(ws, layout)
    return total


# --------------------------------------------------------------------------- scope sheets

def _data_migration(wb: Workbook, spec: dict, layout: TableLayout, sizing: SizingResult,
                    manifest: RenderManifest) -> list[int]:
    """Old Data Migration Scope sheet: one table per delivery wave, each under its wave row."""
    ws = wb.create_sheet(spec["title"])
    last_col = len(layout.columns)
    for col, column in enumerate(layout.columns, start=1):
        ws.column_dimensions[get_column_letter(col)].width = column.width
    waves = sizing.timeline.waves
    tables = sizing.effort.workstreams.data_migration or [[]]
    totals: list[int] = []
    r = 1
    for index, rows in enumerate(tables):
        wave = waves[index] if index < len(waves) else None
        if index:
            r += 2
        if wave is not None:
            values = (wave.name, wave.sequence, wave.partition_axis, wave.total_weeks, wave.hypercare_weeks,
                      ", ".join(wave.countries))
            for row, content, font in ((r, DM_WAVE_HEADERS, TOTAL_FONT), (r + 1, values, None)):
                for col, value in enumerate(content, start=1):
                    cell = ws.cell(row=row, column=col, value=value)
                    cell.fill, cell.border = WAVE_FILL, BORDER
                    if font is not None:
                        cell.font = font
            r += 2

        def band(row: int, text: str, fill: PatternFill) -> None:
            for col in range(1, last_col + 1):
                cell = ws.cell(row=row, column=col)
                cell.fill, cell.border, cell.font = fill, BORDER, TOTAL_FONT
            ws.cell(row=row, column=1, value=text)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)

        band(r, spec["table_title"], TITLE_GREY)
        ws.cell(row=r, column=1).alignment = CENTER
        r += 1
        for col, column in enumerate(layout.columns, start=1):
            cell = ws.cell(row=r, column=col, value=column.header)
            cell.fill, cell.font, cell.alignment, cell.border = NAVY, HEADER_FONT, CENTER, BORDER
        ws.row_dimensions[r].height = 69.5
        header_row = r
        r += 1
        band(r, "Master Data", AMBER)
        r += 1
        table = _Table(ws, layout, manifest, header_row=header_row, first_row=r,
                       wave=wave.name if wave is not None else "", bordered=True)
        for row in rows:
            table.add(f"{row['row_id']}#{index + 1}", {**row, "object": row.get("label") or row["object"]})
        totals.append(table.total_row(bordered=True))
        r = table.row
    _hide_row_ids(ws, layout)
    return totals


def _scope_sheet(wb: Workbook, spec: dict, layout: TableLayout, rows: list[tuple[str, dict]],
                 manifest: RenderManifest, lists: dict[int, tuple[str, ...]] | None = None) -> tuple[Worksheet, int]:
    ws = wb.create_sheet(spec["title"])
    _titled_header(ws, spec["table_title"], layout)
    table = _Table(ws, layout, manifest, header_row=2, first_row=3, bordered=True)
    for row_id, values in rows:
        table.add(row_id, values)
    if rows and lists:
        for col, options in lists.items():
            validation = DataValidation(type="list", formula1=f'"{",".join(options)}"', allow_blank=True)
            ws.add_data_validation(validation)
            letter = get_column_letter(col)
            validation.add(f"{letter}{table.first}:{letter}{table.row - 1}")
    total = table.total_row(bordered=True)
    _hide_row_ids(ws, layout)
    return ws, total


def _analytics(wb: Workbook, spec: dict, layout: TableLayout, sizing: SizingResult, ledger: Ledger,
               manifest: RenderManifest, policy: Policy) -> int:
    """Old Analytics Scope sheet: effort cells are INDEX/MATCH formulas over the editable rate table."""
    cfg = policy.workstreams["analytics"]
    rows = sizing.effort.workstreams.analytics
    ws = wb.create_sheet(spec["title"])
    _titled_header(ws, spec["table_title"], layout)
    levels = list(cfg.get("complexities") or ["Low", "Medium", "High"])
    metrics = list(next(iter(cfg["reference_rates"].values())))
    for letter, width in {"L": 18.0, "M": 13.0, "N": 13.0, "O": 9.0, "P": 9.0, "Q": 9.0}.items():
        ws.column_dimensions[letter].width = width
    title = ws.cell(row=1, column=12, value="Effort Rates (Man Days per Object)")
    title.fill, title.font, title.alignment = TITLE_GREY, TOTAL_FONT, CENTER
    ws.merge_cells(start_row=1, start_column=12, end_row=1, end_column=17)
    for offset, text in enumerate(("Key", "Object Type", "Effort", *levels)):
        cell = ws.cell(row=2, column=12 + offset, value=text)
        cell.fill, cell.font, cell.alignment, cell.border = NAVY, HEADER_FONT, CENTER, BORDER
    r = 3
    for object_type, table in cfg["reference_rates"].items():
        for metric in metrics:
            ws.cell(row=r, column=12, value=f'=M{r}&"|"&N{r}')
            ws.cell(row=r, column=13, value=object_type)
            ws.cell(row=r, column=14, value=metric)
            for offset, rate in enumerate(table[metric]):
                ws.cell(row=r, column=15 + offset, value=rate)
            for col in range(12, 18):
                ws.cell(row=r, column=col).border = BORDER
            r += 1
    last_rate = r - 1
    ws.cell(row=r, column=12, value="Standard content factor").font = TOTAL_FONT
    ws.merge_cells(start_row=r, start_column=12, end_row=r, end_column=14)
    ws.cell(row=r, column=15, value=cfg["standard_factor"]).border = BORDER
    rates, keys, header, factor = f"$O$3:$Q${last_rate}", f"$L$3:$L${last_rate}", "$O$2:$Q$2", f"$O${r}"
    formulas = {}
    for col, metric in zip((7, 8, 9), metrics):
        formulas[col] = (f'=IF($E{{r}}="{IN_SCOPE}",N($D{{r}})*IFERROR(INDEX({rates},MATCH($B{{r}}&"|{metric}",'
                         f'{keys},0),MATCH($F{{r}},{header},0)),0)*IF($C{{r}}="Standard",{factor},1),0)')
    layout = TableLayout(layout.key, layout.section, layout.columns, header_row=2, total_label=layout.total_label,
                         total_cols=layout.total_cols, row_formulas={**formulas, **layout.row_formulas})
    table = _Table(ws, layout, manifest, header_row=2, first_row=3, bordered=True)
    for row in rows:
        table.add(row["row_id"], row)
    if rows:
        for col, options in ((2, tuple(cfg["object_types"])), (3, tuple(cfg["builds"])), (5, (IN_SCOPE, OUT_OF_SCOPE)),
                             (6, tuple(levels))):
            validation = DataValidation(type="list", formula1=f'"{",".join(options)}"', allow_blank=True)
            ws.add_data_validation(validation)
            letter = get_column_letter(col)
            validation.add(f"{letter}{table.first}:{letter}{table.row - 1}")
    total = table.total_row(bordered=True)
    scope = ledger.analytics_scope.data
    if scope is not None and scope.excluded:
        quote = scope.exclusion_evidence[0].quote if scope.exclusion_evidence else ""
        note = ws.cell(row=total + 2, column=1, value=f'Analytics excluded by the RFP: "{quote}"')
        note.font, note.alignment = TOTAL_FONT, WRAP
        ws.merge_cells(start_row=total + 2, start_column=1, end_row=total + 2, end_column=len(layout.columns))
    _hide_row_ids(ws, layout)
    return total


# --------------------------------------------------------------------------- Project Timeline

def _project_timeline(ws: Worksheet, sizing: SizingResult, manifest: RenderManifest, pes: str,
                      policy: Policy) -> None:
    """Old Resources sheet ('Project Timeline'): one Skill x month grid per wave (onsite rows, Onsite
    Total, offshore rows, Offshore Total, Grand Total), the combined programme calendar for several
    waves, and the reconciliation with the Summary of Project Effort."""
    plan, days = sizing.plan, policy.commercials.working_days_per_month
    title = ws.cell(row=1, column=1, value="Resource Estimation")
    title.font, title.fill = HEADER_FONT, HEADER_FILL
    ws.cell(row=1, column=2, value=(f"source: {plan.source} (derived = deterministic); 1 man-month = {days} "
                                    "person-days; FTE figures are quoted in steps of 0.5")).fill = HEADER_FILL
    for col in (3, 4):
        ws.cell(row=1, column=col).fill = HEADER_FILL
    r = 1
    grand_rows: list[tuple[Any, int]] = []
    widest = max((g.months for g in plan.grids), default=0)
    rid_col = GRID_FIRST_MONTH_COL + max(widest, 1) + 1
    for grid in plan.grids:
        r += 1
        header = ws.cell(row=r, column=1, value=f"Resource Plan - {grid.wave_name}")
        header.font, header.fill = TOTAL_FONT, GRID_FILL
        for col in range(2, grid.months + 4):
            ws.cell(row=r, column=col).fill = GRID_FILL
        r += 1
        ws.cell(row=r, column=1, value="Skill").font = TOTAL_FONT
        ws.cell(row=r, column=2, value="Location").font = TOTAL_FONT
        for i in range(grid.months):
            ws.cell(row=r, column=GRID_FIRST_MONTH_COL + i, value=f"M{i + 1}").font = TOTAL_FONT
        total_col = GRID_FIRST_MONTH_COL + grid.months
        ws.cell(row=r, column=total_col, value="Total Man-months").font = TOTAL_FONT
        header_row = r
        r += 1
        ws.cell(row=r, column=1, value="Phase")
        for i, band in enumerate(grid.phase_bands):
            ws.cell(row=r, column=GRID_FIRST_MONTH_COL + i, value=band)
        r += 1
        first_l, last_l = get_column_letter(GRID_FIRST_MONTH_COL), get_column_letter(total_col - 1)
        columns = {1: ColumnRef(field="role_title", kind="label"), 2: ColumnRef(field="location", kind="label"),
                   **{GRID_FIRST_MONTH_COL + i: ColumnRef(field=f"M{i + 1}", kind=INPUT) for i in range(grid.months)},
                   total_col: ColumnRef(field="total", kind=FORMULA)}
        entry = TableManifest(key="grid", section="grid", sheet=ws.title, header_row=header_row, label_col=1,
                              total_label=GRID_TOTAL_LABEL, row_id_col=rid_col, columns=columns,
                              wave=grid.wave_name, skip_labels=list(GRID_SKIP_LABELS))
        manifest.tables.append(entry)

        def write_row(label: str, location: str, cells: list[Any], bold: bool = False) -> int:
            nonlocal r
            ws.cell(row=r, column=1, value=label).font = TOTAL_FONT if bold else Font()
            ws.cell(row=r, column=2, value=location)
            for i, value in enumerate(cells):
                cell = ws.cell(row=r, column=GRID_FIRST_MONTH_COL + i, value=value)
                if bold:
                    cell.font = TOTAL_FONT
            total_cell = ws.cell(row=r, column=total_col, value=f"=SUM({first_l}{r}:{last_l}{r})")
            if bold:
                total_cell.font = TOTAL_FONT
            r += 1
            return r - 1

        subtotal: dict[str, int] = {}
        for location, label in (("Onsite", "Onsite Total"), ("Offshore", "Offshore Total")):
            first = r
            for row in grid.rows:
                if row.location != location:
                    continue
                values = [round(row.fte_at(i + 1), 2) or None for i in range(grid.months)]
                at = write_row(row.role_title, row.location, values)
                row_id = f"{grid.wave_name}||{row.role_title}||{row.location}"
                ws.cell(row=at, column=rid_col, value=row_id)
                entry.rows[row_id] = {f"M{i + 1}": v for i, v in enumerate(values)}
            last = r - 1
            sums = [(f"=SUM({get_column_letter(GRID_FIRST_MONTH_COL + i)}{first}:"
                     f"{get_column_letter(GRID_FIRST_MONTH_COL + i)}{last})" if last >= first else None)
                    for i in range(grid.months)]
            subtotal[location] = write_row(label, location, sums, bold=True)
        on, off = subtotal["Onsite"], subtotal["Offshore"]
        grand = write_row(GRID_TOTAL_LABEL, "", [f"={get_column_letter(GRID_FIRST_MONTH_COL + i)}{on}+"
                                                 f"{get_column_letter(GRID_FIRST_MONTH_COL + i)}{off}"
                                                 for i in range(grid.months)], bold=True)
        grand_rows.append((grid, grand))

    if len(plan.grids) > 1:
        combined = plan.combined_by_month
        r += 1
        header = ws.cell(row=r, column=1, value="Combined Programme Staffing")
        header.font, header.fill = TOTAL_FONT, GRID_FILL
        ws.cell(row=r, column=2, value=(f"all waves on one calendar ({sizing.timeline.sequencing}); "
                                        f"peak {plan.peak_fte:g} FTE"))
        for col in range(3, len(combined) + 4):
            ws.cell(row=r, column=col).fill = GRID_FILL
        r += 1
        ws.cell(row=r, column=1, value="Month").font = TOTAL_FONT
        for i in range(len(combined)):
            ws.cell(row=r, column=GRID_FIRST_MONTH_COL + i, value=f"M{i + 1}").font = TOTAL_FONT
        ws.cell(row=r, column=GRID_FIRST_MONTH_COL + len(combined), value="Total Man-months").font = TOTAL_FONT
        r += 1
        ws.cell(row=r, column=1, value="Total FTE").font = TOTAL_FONT
        ws.cell(row=r, column=2, value="all waves")
        for i in range(len(combined)):
            parts = [f"{get_column_letter(GRID_FIRST_MONTH_COL + i - g.start_month)}{row}"
                     for g, row in grand_rows if g.start_month <= i < g.start_month + g.months]
            ws.cell(row=r, column=GRID_FIRST_MONTH_COL + i, value=("=" + "+".join(parts)) if parts else None).font = TOTAL_FONT
        ws.cell(row=r, column=GRID_FIRST_MONTH_COL + len(combined),
                value=f"=SUM(C{r}:{get_column_letter(2 + len(combined))}{r})").font = TOTAL_FONT
        r += 1
        ws.cell(row=r, column=1, value="Waves Active")
        for i in range(len(combined)):
            active = sum(1 for g in plan.grids if g.start_month <= i < g.start_month + g.months)
            ws.cell(row=r, column=GRID_FIRST_MONTH_COL + i, value=active or None)
        r += 1

    # Reconciliation with the Summary of Project Effort (every figure a live link to it).
    ladder = policy.commercials.summary_ladder
    r += 1
    header = ws.cell(row=r, column=1, value="Reconciliation with Summary of Project Effort")
    header.font, header.fill = TOTAL_FONT, GRID_FILL
    for col in range(2, 5):
        ws.cell(row=r, column=col).fill = GRID_FILL
    r += 1
    ws.cell(row=r, column=1, value="Component").font = TOTAL_FONT
    ws.cell(row=r, column=2, value="Person-Days").font = TOTAL_FONT
    r += 1

    def line(label: str, value: Any, bold: bool = False) -> int:
        nonlocal r
        ws.cell(row=r, column=1, value=label).font = TOTAL_FONT if bold else Font()
        ws.cell(row=r, column=2, value=value).font = TOTAL_FONT if bold else Font()
        r += 1
        return r - 1

    line("Total Build Effort (Functional + Tech Dev + ...)", f"={pes}!B9", True)
    line(f"Plus Final Preparation ({ladder.final_prep_pct:g}% of Build)", f"={pes}!B10")
    line(f"Plus Go-Live Effort ({ladder.go_live_pct:g}% of Build)", f"={pes}!B11")
    line(f"Plus Programme Management ({ladder.project_mgmt_pct:g}% of Build)", f"={pes}!B12")
    line("Plus Hypercare", f"={pes}!B13")
    line("Sum: Summary of Imp", f"={pes}!B14", True)
    line(f"Plus Risk Contingency ({round((ladder.risk_factor - 1) * 100):g}% of Summary of Imp)", f"={pes}!B15")
    line("Sum: Total Project Effort", f"={pes}!B16", True)
    if grand_rows:
        cells = "+".join(f"{get_column_letter(GRID_FIRST_MONTH_COL + g.months)}{row}" for g, row in grand_rows)
        grid_row = line("Grid rows sum (all waves x days-per-month)", f"=ROUND(({cells})*{days},2)")
        expected = line("Expected grid rows sum (= Summary of Imp)", f"={pes}!B14")
        line("Variance", f"=ROUND(B{grid_row}-B{expected},2)")
        if len(grand_rows) > 1:
            r += 1
            for col, text in enumerate(("Per Wave", "Grid (PD)", "Attributed scope (PD)", "Source"), start=1):
                ws.cell(row=r, column=col, value=text).font = TOTAL_FONT
            r += 1
            we = sizing.wave_effort
            for i, (g, row) in enumerate(grand_rows):
                ws.cell(row=r, column=1, value=g.wave_name)
                ws.cell(row=r, column=2, value=f"=ROUND({get_column_letter(GRID_FIRST_MONTH_COL + g.months)}{row}*{days},2)")
                if i < len(we.wave_names):
                    ws.cell(row=r, column=3, value=we.wave_total(i))
                    ws.cell(row=r, column=4, value=we.source.get(g.wave_name, ""))
                r += 1
        for flag in sizing.wave_effort.flags:
            ws.cell(row=r, column=1, value=f"Assumption: {flag}").font = Font(italic=True, size=9)
            r += 1
    _autosize(ws)
    _hide_row_ids(ws, rid_col)


# --------------------------------------------------------------------------- Summary of Project Effort

def _summary(ws: Worksheet, links: dict[str, str | None], sizing: SizingResult, manifest: RenderManifest,
             policy: Policy) -> None:
    """Old Summary of Project Effort (rows 1-16, labels verbatim incl. the template's typos), with the
    two rates added to the side table and a Total Project Cost row (17)."""
    s = sizing.summary
    ws.cell(row=1, column=1, value="Summary Project Effort")
    ws.cell(row=2, column=1, value="Scope Item")
    ws.cell(row=2, column=2, value="Man Days Effort")
    params = {}
    for offset, (label, key) in enumerate(SUMMARY_PARAMETERS):
        row = 2 + offset
        ws.cell(row=row, column=4, value=label)
        ws.cell(row=row, column=5, value=getattr(s, key))
        manifest.settings[key] = getattr(s, key)
        params[key] = f"$E${row}"
    fallback = {"functional_scope": s.functional_scope, "data_migration": s.data_migration, "tech_dev": s.tech_dev,
                "security": s.security, "basis": s.basis, "analytics": s.analytics}
    for spec in policy.workbook["project_summary_rows"]:
        ws.cell(row=spec["row"], column=1, value=spec["label"])
        ws.cell(row=spec["row"], column=2, value=links.get(spec["source"]) or fallback[spec["source"]])
    rows = [
        (9, "Total Build Effort(Upto Realization Phase)", "=SUM(B3:B8)", True),
        (10, "Final preparation phase", f"=ROUND(B9*{params['final_prep_pct']}/100,2)", False),
        (11, "Go-Live Effort", f"=ROUND(B9*{params['go_live_pct']}/100,2)", False),
        (12, "Project Mangement Effort", f"=ROUND(B9*{params['project_mgmt_pct']}/100,2)", False),
        (13, HYPERCARE_LABEL, s.hypercare, False),
        (14, "Summary of Imp", "=SUM(B9:B13)", True),
        (15, "Risk Contingency", f"=ROUND(B14*({params['risk_factor']}-1),2)", True),
        (16, "Total Project Effort", "=B14+B15", True),
        (17, "Total Project Cost (USD)",
         f"=ROUND((B16-B13)*{params['daily_rate_usd']}+B13*{params['hypercare_rate_usd']},2)", True),
    ]
    manifest.settings[HYPERCARE_SETTING] = s.hypercare
    for row, label, value, bold in rows:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
        if bold:
            ws.cell(row=row, column=1).font = ws.cell(row=row, column=2).font = TOTAL_FONT
    for row in ws.iter_rows(min_row=1, max_row=17, min_col=1, max_col=2):
        for cell in row:
            cell.border = BORDER
    for row in ws.iter_rows(min_row=2, max_row=1 + len(SUMMARY_PARAMETERS), min_col=4, max_col=5):
        for cell in row:
            cell.border = BORDER
    ws.merge_cells("A1:B1")
    ws["A1"].fill, ws["A1"].font, ws["A1"].alignment = AMBER, Font(bold=True), Alignment(horizontal="center")
    for row in range(2, 2 + len(SUMMARY_PARAMETERS)):
        ws.cell(row=row, column=4).fill, ws.cell(row=row, column=4).font = NAVY, HEADER_FONT
    for coord in ("A2", "B2"):
        ws[coord].fill, ws[coord].font, ws[coord].alignment = NAVY, HEADER_FONT, Alignment(horizontal="center")
    ws["E5"].font = Font(bold=True)
    for total_row in (9, 16, 17):
        for col in (1, 2):
            ws.cell(row=total_row, column=col).fill = GREEN
    _autosize(ws)


def _bid_sheet(ws: Worksheet, manifest: RenderManifest, ledger: Ledger) -> None:
    ws.sheet_state = "hidden"
    rows = [("bid_id", manifest.bid_id), ("render_id", manifest.render_id),
            ("ledger_version", manifest.ledger_version), ("policy_version", manifest.policy_version),
            ("rendered_at", manifest.rendered_at), ("client", ledger.meta.client_name),
            ("rate_card_sheet", ledger.meta.rate_card_sheet),
            ("note", "Do not edit or delete this sheet: it links the workbook to its bid for the proposal step.")]
    for r, (key, value) in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=key)
        ws.cell(row=r, column=2, value=value)


def _lob_labels(sizing: SizingResult) -> dict[str, str]:
    """Old _summary_lob_label: the LOB with its cross-mapped module names in brackets."""
    return {lob: f"{lob} [{', '.join(m.submodules)}]" if m.submodules else lob
            for lob, m in sizing.effort.modules.items()}


def render_workbook(ledger: Ledger, sizing: SizingResult, path: Path, policy: Policy | None = None,
                    manifest_dir: Path | None = None) -> Path:
    """Write the effort workbook to `path`; the render manifest goes to `manifest_dir` when given."""
    policy = policy or get_policy()
    spec = policy.workbook["sheets"]
    layouts = table_layouts(policy.workbook)
    manifest = RenderManifest(render_id=uuid.uuid4().hex[:12], bid_id=ledger.meta.bid_id,
                              ledger_version=ledger.meta.version, policy_version=policy.version,
                              rendered_at=datetime.now(timezone.utc).isoformat(timespec="seconds"),
                              file_name=path.name, summary_sheet=spec["project_summary"]["title"])
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = spec["project_summary"]["title"]
    functional_ws = wb.create_sheet(spec["functional_scope"]["title"])
    timeline_ws = wb.create_sheet(spec["project_timeline"]["title"])
    used = {ws.title for ws in wb.worksheets} | {spec[k]["title"] for k in (
        "non_catalogue", "tech_dev", "data_migration", "basis", "security", "analytics", "bid")}
    modules = _module_sheets(wb, layouts["catalogue"], sizing, manifest, used)
    functional_total = _functional_scope(functional_ws, spec["functional_scope"]["headers"], modules, sizing,
                                         _lob_labels(sizing))
    nc_total = _non_catalogue(wb, spec["non_catalogue"]["title"], layouts["non_catalogue"], sizing, manifest)
    td_total = _tech_dev(wb, spec["tech_dev"]["title"], layouts["tech_dev"], sizing, manifest)
    dm_totals = _data_migration(wb, spec["data_migration"], layouts["data_migration"], sizing, manifest)
    tables = sizing.effort.workstreams
    _, basis_total = _scope_sheet(wb, spec["basis"], layouts["basis"], [(r["row_id"], r) for r in tables.basis],
                                  manifest)
    _, sec_total = _scope_sheet(wb, spec["security"], layouts["security"],
                                [(r["row_id"], r) for r in tables.security], manifest,
                                lists={2: (IN_SCOPE, OUT_OF_SCOPE), 3: ("Low", "Medium", "High")})
    ana_total = _analytics(wb, spec["analytics"], layouts["analytics"], sizing, ledger, manifest, policy)

    def link(sheet: str, col: str, rows: list[int | None]) -> str | None:
        rows = [r for r in rows if r]
        return "=" + "+".join(f"{quote_sheet(sheet)}!{col}{r}" for r in rows) if rows else None

    functional_parts = [x[1:] for x in (link(functional_ws.title, "D", [functional_total]),
                                        link(spec["non_catalogue"]["title"], "D", [nc_total])) if x]
    links = {"functional_scope": ("=" + "+".join(functional_parts)) if functional_parts else None,
             "data_migration": link(spec["data_migration"]["title"], "I", dm_totals),
             "tech_dev": link(spec["tech_dev"]["title"], "M", [td_total]),
             "security": link(spec["security"]["title"], "E", [sec_total]),
             "basis": link(spec["basis"]["title"], "F", [basis_total]),
             "analytics": link(spec["analytics"]["title"], "J", [ana_total])}
    _summary(summary_ws, links, sizing, manifest, policy)
    _project_timeline(timeline_ws, sizing, manifest, quote_sheet(summary_ws.title), policy)
    _bid_sheet(wb.create_sheet(spec["bid"]["title"]), manifest, ledger)

    head = [summary_ws.title, functional_ws.title, timeline_ws.title]
    tail = [spec[k]["title"] for k in ("non_catalogue", "tech_dev", "data_migration", "basis", "security",
                                       "analytics", "bid")]
    names = wb.sheetnames
    order = [n for n in head if n in names] + [n for n in names if n not in head and n not in tail] \
        + [n for n in tail if n in names]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    if manifest_dir is not None:
        save_manifest(manifest, manifest_dir)
    return path


__all__ = ["BID_SHEET", "render_workbook"]

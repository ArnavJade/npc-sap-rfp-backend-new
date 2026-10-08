"""Render manifest and the reviewer-edit diff (call 2).

Call 1 records, per render, every table row it wrote (by row id) with its input values in
`ledger/manifest-<render_id>.json`, and stamps the render id in the workbook's hidden `_bid` sheet.
Call 2 reads the reviewed workbook back against that manifest:

* rows are matched by the hidden row-id column, never by position, so inserted/deleted/re-sorted rows
  are handled; a table ends at its 'Total' label (or the next table title);
* a changed `input` cell -> an edit; a row id that disappeared -> deleted_row; a filled row with no
  row id inside a table -> added_row; a `formula` cell overwritten with a constant, or a changed
  `derived` cell -> conflict (reported, not applied);
* Summary of Project Effort parameters and the Hypercare row are found by their labels -> settings.

`apply_overrides(ledger, edits)` then writes them into the ledger: scope-sheet edits (non-catalogue,
data migration, basis, security, analytics) change the ledger rows directly; catalogue lines, Tech Dev
lines, grid cells and settings become applied `ledger.overrides` that sizing honours.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel, Field

from bidcore.ledger.models import Ledger
from bidcore.ledger.sections_effort import (
    AnalyticsObject, BasisActivity, DataMigrationObject, NonCatalogueItem, ScopeItem, SecurityActivity,
)
from bidcore.ledger.sections_proposal import Override
from bidcore.render.workbook.layout import (
    DERIVED, FORMULA, HYPERCARE_LABEL, HYPERCARE_SETTING, INPUT, LABEL, SUMMARY_PARAMETERS,
)

BID_SHEET = "_bid"
_TITLE_PREFIXES = ("Data Migration - ", "Resource Plan - ")


class ColumnRef(BaseModel):
    field: str
    kind: str


class TableManifest(BaseModel):
    key: str
    section: str
    sheet: str
    header_row: int
    label_col: int
    total_label: str = "Total"
    row_id_col: int
    columns: dict[int, ColumnRef]
    rows: dict[str, dict[str, Any]] = Field(default_factory=dict)
    wave: str = ""
    skip_labels: list[str] = Field(default_factory=list)   # band / subtotal rows inside the table


class RenderManifest(BaseModel):
    render_id: str
    bid_id: str
    ledger_version: int
    policy_version: str
    rendered_at: str
    file_name: str = ""
    summary_sheet: str = ""
    settings: dict[str, float] = Field(default_factory=dict)
    lob_sheets: dict[str, str] = Field(default_factory=dict)      # sheet title -> LOB
    tables: list[TableManifest] = Field(default_factory=list)


class WorkbookError(ValueError):
    """The upload is not a workbook this service produced (no `_bid` sheet / unknown render)."""


def manifest_path(ledger_dir: Path, render_id: str) -> Path:
    return ledger_dir / f"manifest-{render_id}.json"


def save_manifest(manifest: RenderManifest, ledger_dir: Path) -> Path:
    ledger_dir.mkdir(parents=True, exist_ok=True)
    path = manifest_path(ledger_dir, manifest.render_id)
    path.write_text(manifest.model_dump_json(indent=1), encoding="utf-8")
    return path


def read_bid_sheet(workbook: Path) -> dict[str, str]:
    """{bid_id, render_id, ...} from the hidden `_bid` sheet; WorkbookError when absent."""
    wb = load_workbook(workbook, read_only=True)
    try:
        if BID_SHEET not in wb.sheetnames:
            raise WorkbookError("this workbook has no '_bid' sheet: it was not produced by this service. "
                                "Regenerate the effort workbook (call 1) and review that one.")
        out = {}
        for row in wb[BID_SHEET].iter_rows(min_row=1, max_col=2, values_only=True):
            if row and row[0]:
                out[str(row[0]).strip()] = "" if row[1] is None else str(row[1]).strip()
        if not out.get("bid_id") or not out.get("render_id"):
            raise WorkbookError("the '_bid' sheet is incomplete (bid_id / render_id missing)")
        return out
    finally:
        wb.close()


# --------------------------------------------------------------------------- diff

def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _norm(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return round(float(value), 6)
    text = str(value).strip()
    if not text:
        return None
    try:
        return round(float(text), 6)
    except ValueError:
        return text


def _same(a: Any, b: Any) -> bool:
    a, b = _norm(a), _norm(b)
    if isinstance(a, float) and isinstance(b, float):
        return abs(a - b) < 1e-6
    return a == b


def _ends_table(text: Any, total_label: str) -> bool:
    label = str(text or "").strip()
    return label.lower().startswith(total_label.lower()) or label.startswith(_TITLE_PREFIXES)


def _diff_table(ws, table: TableManifest) -> list[Override]:
    edits: list[Override] = []
    seen: set[str] = set()
    blank_run = 0
    row = table.header_row + 1
    while row <= ws.max_row and blank_run < 3:
        label = ws.cell(row=row, column=table.label_col).value
        if _ends_table(label, table.total_label):
            total_col = next((c for c, ref in table.columns.items() if ref.field == "total"), None)
            if total_col and not _is_formula(ws.cell(row=row, column=total_col).value):
                edits.append(Override(sheet=table.sheet, cell=ws.cell(row=row, column=total_col).coordinate,
                                      section=table.section, field="total", kind="conflict",
                                      new=ws.cell(row=row, column=total_col).value,
                                      note="the table Total was overwritten with a constant; it is recalculated"))
            break
        values = {c: ws.cell(row=row, column=c).value for c in table.columns}
        row_id = str(ws.cell(row=row, column=table.row_id_col).value or "").strip()
        first_text = next((str(v).strip() for v in values.values() if v is not None and str(v).strip()), "")
        if not row_id and (str(label or "").strip() in table.skip_labels or first_text in table.skip_labels):
            row += 1
            continue
        if not row_id and all(v is None or str(v).strip() == "" for v in values.values()):
            blank_run += 1
            row += 1
            continue
        blank_run = 0
        if row_id and row_id in table.rows:
            seen.add(row_id)
            old = table.rows[row_id]
            for col, ref in table.columns.items():
                cell = ws.cell(row=row, column=col)
                if ref.kind == FORMULA:
                    if not _is_formula(cell.value) and cell.value is not None:
                        edits.append(Override(sheet=table.sheet, cell=cell.coordinate, section=table.section,
                                              row_id=row_id, field=ref.field, new=cell.value, kind="conflict",
                                              note="formula overwritten with a constant; edit the input cells instead"))
                    continue
                if ref.kind == LABEL or _same(cell.value, old.get(ref.field)):
                    continue
                kind = "edit" if ref.kind == INPUT else "conflict"
                note = "" if kind == "edit" else "computed cell changed; edit the inputs (count, type, complexity)"
                edits.append(Override(sheet=table.sheet, cell=cell.coordinate, section=table.section, row_id=row_id,
                                      field=ref.field, old=old.get(ref.field), new=_norm(cell.value), kind=kind,
                                      note=note or (f"wave:{table.wave}" if table.wave else "")))
        elif not row_id:
            data = {ref.field: _norm(values[c]) for c, ref in table.columns.items()
                    if ref.kind != FORMULA and values[c] is not None}
            edits.append(Override(sheet=table.sheet, cell=f"A{row}", section=table.section, kind="added_row",
                                  new=data, note=f"wave:{table.wave}" if table.wave else ""))
        row += 1
    for row_id in table.rows:
        if row_id not in seen:
            edits.append(Override(sheet=table.sheet, section=table.section, row_id=row_id, kind="deleted_row",
                                  note=f"wave:{table.wave}" if table.wave else ""))
    return edits


def _find_label(ws, column: int, label: str) -> int | None:
    for row in range(1, ws.max_row + 1):
        if str(ws.cell(row=row, column=column).value or "").strip() == label:
            return row
    return None


def _diff_settings(ws, manifest: RenderManifest) -> list[Override]:
    edits: list[Override] = []
    targets = [(4, label, key) for label, key in SUMMARY_PARAMETERS] + [(1, HYPERCARE_LABEL, HYPERCARE_SETTING)]
    for label_col, label, key in targets:
        row = _find_label(ws, label_col, label)
        if row is None or key not in manifest.settings:
            continue
        cell = ws.cell(row=row, column=label_col + 1)
        if _is_formula(cell.value) or _same(cell.value, manifest.settings[key]):
            continue
        value = _norm(cell.value)
        kind = "setting" if isinstance(value, float) and value >= 0 else "conflict"
        edits.append(Override(sheet=ws.title, cell=cell.coordinate, field=key, old=manifest.settings[key],
                              new=value, kind=kind, note="" if kind == "setting" else "not a non-negative number"))
    return edits


def read_edits(workbook: Path, ledger_dir: Path) -> tuple[RenderManifest, list[Override]]:
    """The reviewer's edits in `workbook` relative to the render it came from."""
    stamp = read_bid_sheet(workbook)
    path = manifest_path(ledger_dir, stamp["render_id"])
    if not path.is_file():
        raise WorkbookError(f"render {stamp['render_id']} of bid {stamp['bid_id']} is unknown here "
                            "(manifest missing) - was the bid created on another server?")
    manifest = RenderManifest.model_validate_json(path.read_text(encoding="utf-8"))
    wb = load_workbook(workbook)            # formulas, not cached values
    try:
        edits: list[Override] = []
        for table in manifest.tables:
            if table.sheet not in wb.sheetnames:
                edits += [Override(sheet=table.sheet, section=table.section, row_id=rid, kind="deleted_row",
                                   note="sheet removed") for rid in table.rows]
                continue
            edits += _diff_table(wb[table.sheet], table)
        if manifest.summary_sheet in wb.sheetnames:
            edits += _diff_settings(wb[manifest.summary_sheet], manifest)
        return manifest, _collapse_dm_deletions(manifest, edits)
    finally:
        wb.close()


def _collapse_dm_deletions(manifest: RenderManifest, edits: list[Override]) -> list[Override]:
    """A conversion object is one ledger row shown in every wave table: it is deleted only when the
    reviewer removed it from all of them; a partial removal is reported as a conflict."""
    tables = sum(1 for t in manifest.tables if t.key == "data_migration") or 1
    deleted: dict[str, list[Override]] = {}
    out: list[Override] = []
    for edit in edits:
        if edit.section == "data_migration" and edit.kind == "deleted_row":
            deleted.setdefault(_base_row_id(edit.section, edit.row_id), []).append(edit)
        else:
            out.append(edit)
    for base, group in deleted.items():
        if len(group) >= tables:
            out.append(Override(sheet=group[0].sheet, section="data_migration", row_id=base, kind="deleted_row"))
        else:
            out += [e.model_copy(update={"kind": "conflict", "note": (e.note + "; " if e.note else "")
                                         + "deleted from some wave tables only; delete it from every wave to drop it"})
                    for e in group]
    return out


# --------------------------------------------------------------------------- apply

_LEDGER_ROW_MODELS = {"non_catalogue": NonCatalogueItem, "data_migration": DataMigrationObject,
                      "basis": BasisActivity, "security": SecurityActivity, "analytics": AnalyticsObject}
_OVERRIDE_SECTIONS = {"scope_items", "tech_dev", "grid"}
_NUMERIC_INT = {"dev", "qa", "prd", "effort_days", "no_of_objects"}


def _base_row_id(section: str, row_id: str) -> str:
    return row_id.split("#", 1)[0] if section == "data_migration" else row_id


_TEXT_FIELDS = {"status", "complexity", "object_type", "build", "object", "activity", "name", "description",
                "sap_product", "category", "sap_module"}


def _coerce(field: str, value: Any) -> Any:
    if value is None:
        return None
    if field in _NUMERIC_INT and isinstance(value, (int, float)):
        return int(round(value))
    if field in _TEXT_FIELDS:
        return str(value).strip()
    return value


def _next_row_id(rows: list, prefix: str) -> str:
    used = {r.row_id for r in rows}
    n = len(rows) + 1
    while f"{prefix}-{n}" in used:
        n += 1
    return f"{prefix}-{n}"


def _apply_ledger_edit(ledger: Ledger, edit: Override) -> str:
    """Direct ledger change for scope-sheet rows; returns '' or a reason it was not applied."""
    section = getattr(ledger, edit.section)
    model = _LEDGER_ROW_MODELS[edit.section]
    prefix = {"non_catalogue": "nc", "data_migration": "dm", "basis": "bas", "security": "sec",
              "analytics": "ana"}[edit.section]
    if edit.kind == "added_row":
        data = {k: v for k, v in dict(edit.new or {}).items() if k in model.model_fields and v is not None}
        if edit.section == "non_catalogue" and "name" not in data:
            return "added row has no SAP tool name"
        key = {"data_migration": "object", "basis": "activity", "security": "activity",
               "analytics": "object"}.get(edit.section)
        if key and key not in data:
            return f"added row has no {key}"
        data = {k: _coerce(k, v) for k, v in data.items()}
        try:
            row = model(**data, row_id=_next_row_id(section.rows, prefix), written_by="reviewer",
                        note="added in the reviewed workbook")
        except Exception as exc:   # invalid enum etc.
            return f"added row rejected: {exc}"[:200]
        if edit.section == "data_migration" and any(r.object.strip().lower() == row.object.strip().lower()
                                                    for r in section.rows):
            return "object already listed (added in another wave table)"
        section.rows.append(row)
        section.state = "written"
        edit.row_id = row.row_id
        return ""
    base = _base_row_id(edit.section, edit.row_id)
    target = next((r for r in section.rows if r.row_id == base), None)
    if target is None:
        return "row no longer in the ledger"
    if edit.kind == "deleted_row":
        section.rows.remove(target)
        return ""
    if edit.field not in model.model_fields:
        return f"field {edit.field} is not editable"
    try:
        updated = model.model_validate({**target.model_dump(), edit.field: _coerce(edit.field, edit.new)})
    except Exception as exc:
        return f"value rejected: {exc}"[:200]
    section.rows[section.rows.index(target)] = updated
    return ""


def _add_scope_item(ledger: Ledger, edit: Override) -> str:
    data = dict(edit.new or {})
    scope_id = str(data.get("scope_id") or "").strip()
    if not scope_id:
        return "added catalogue line has no Scope ID"
    country = str(data.get("country") or "").strip().upper()
    item = ScopeItem(row_id=_next_row_id(ledger.scope_items.rows, "si"), scope_item_id=scope_id,
                     lob=str(data.get("lob") or ""), business_area=str(data.get("business_area") or ""),
                     description=str(data.get("description") or ""), countries=[country] if country else [],
                     mapping_basis="reviewer", written_by="reviewer", note="added in the reviewed workbook")
    ledger.scope_items.rows.append(item)
    line = f"{item.row_id}@{country or 'ALL'}"
    for field in ("workshops_configuration", "unit_testing", "integration_testing", "documentation_training",
                  "user_acceptance_testing", "multiplication_factor", "wave"):
        if data.get(field) is not None:
            ledger.overrides.append(Override(sheet=edit.sheet, section="scope_items", row_id=line, field=field,
                                             new=data[field], kind="edit", applied=True, note="added line"))
    edit.row_id = line
    return ""


def apply_overrides(ledger: Ledger, edits: list[Override]) -> str:
    """Write the reviewer's edits into the ledger (in place); returns a one-line summary."""
    for edit in edits:
        # A later wave's Data Migration table is the base table scaled for that wave: an edit there
        # changes that wave only (sizing applies it); wave 1's table is the base itself.
        wave_no = (edit.row_id or "").rpartition("#")[2] if "#" in (edit.row_id or "") else ""
        if edit.section == "data_migration" and edit.kind == "edit" and wave_no not in ("", "1"):
            edit.section = "data_migration_wave"
    replaced = {(e.section, e.row_id, e.field) for e in edits if e.kind in ("edit", "setting")}
    ledger.overrides = [o for o in ledger.overrides if (o.section, o.row_id, o.field) not in replaced]
    counts = {"applied": 0, "skipped": 0, "conflicts": 0}
    for edit in edits:
        reason = ""
        if edit.kind == "conflict":
            counts["conflicts"] += 1
            ledger.overrides.append(edit)
            continue
        if edit.kind == "setting":
            pass
        elif edit.section == "data_migration_wave":
            pass                                   # applied by sizing to that wave's table
        elif edit.section in _OVERRIDE_SECTIONS:
            if edit.section == "scope_items" and edit.kind == "added_row":
                reason = _add_scope_item(ledger, edit)
        elif edit.section in _LEDGER_ROW_MODELS:
            reason = _apply_ledger_edit(ledger, edit)
        else:
            reason = f"unknown section '{edit.section}'"
        edit.applied = not reason
        if reason:
            edit.note = (edit.note + "; " if edit.note else "") + reason
        counts["applied" if edit.applied else "skipped"] += 1
        ledger.overrides.append(edit)
    return (f"{counts['applied']} edit(s) applied, {counts['skipped']} skipped, "
            f"{counts['conflicts']} conflict(s)")


def edits_report(edits: list[Override]) -> str:
    """Markdown list of the edits for notes/ and the job record."""
    if not edits:
        return "No reviewer edits found in the workbook."
    lines = ["| Sheet | Cell | Row | Field | Old | New | Kind | Applied | Note |", "|---|---|---|---|---|---|---|---|---|"]
    for e in edits:
        new = json.dumps(e.new, default=str) if isinstance(e.new, dict) else e.new
        lines.append(f"| {e.sheet} | {e.cell} | {e.row_id} | {e.field} | {e.old} | {new} | {e.kind} | "
                     f"{'yes' if e.applied else 'no'} | {e.note} |")
    return "\n".join(lines)

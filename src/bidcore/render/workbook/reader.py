"""Call 2 from ANY effort workbook in the template layout - no `_bid` sheet, no call-1 workspace.

The reviewed workbook is the single source of the numbers: every sheet of the effort-estimation
template is read back by its header texts (not by sheet titles or a hidden manifest), so a workbook
produced by the old agent, edited by hand, or produced on another server works the same way:

    read_effort_workbook(path) -> WorkbookSnapshot
        .sizing(policy)  -> SizingResult    every figure / table / diagram of the Word document
        .ledger(bid_id)  -> Ledger          the facts the section writers draft from

Inputs are taken as typed (constants); formula cells are recomputed exactly as the workbook would
(row Totals, the Summary of Project Effort ladder), so a workbook saved by openpyxl - which has no
cached formula values - reads the same as one saved by Excel. Where a value cell holds a formula,
its cached value is used when there is one.

Read per sheet (identified by its header row):
  Summary of Project Effort   ladder parameters (D/E), Hypercare row
  Functional Scope Estimation sub-module names per LOB ("Finance [Finance, Controlling (CO)]")
  Project Timeline            one resource grid per "Resource Plan - <wave>" block
  <LOB> sheets                catalogue lines (header "Line of Business")
  Non Catalogue SAP Tools     tool rows
  Tech Dev Scope              RICEFW / Fiori / interface / third-party rows
  Data Migration Scope        one conversion-object table per wave, with the wave's duration row
  Basis / Security / Analytics Scope
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel, Field

from bidcore.effort.models import EffortModel, EffortRow, ModuleEffort, NonCatalogueEffort, TechDevRow, WorkstreamTables
from bidcore.effort.summary import compute_summary
from bidcore.effort.waves import WaveEffort
from bidcore.effort.workstreams import analytics_object_effort, workstream_totals
from bidcore.ledger.base import Evidence
from bidcore.ledger.models import Ledger, new_ledger
from bidcore.ledger.sections_effort import (
    AnalyticsObject, BasisActivity, CountryScope, DataMigrationObject, Integration, NonCatalogueItem, RfpProfile,
    ScopeItem, SecurityActivity, Timeline, Wave,
)
from bidcore.policy import Policy, get_policy
from bidcore.render.workbook.layout import HYPERCARE_LABEL, SUMMARY_PARAMETERS
from bidcore.resourcing.models import OFFSHORE, ONSITE, GridCell, GridRow, PassWeights, ResourcePlan, WaveResourceGrid
from bidcore.timeline.model import ResolvedTimeline, ResolvedWave, weeks_for

PHASES = ("workshops_configuration", "unit_testing", "integration_testing", "documentation_training",
          "user_acceptance_testing")
DM_KEYS = ("func_spec", "program_dev", "iteration_1", "iteration_2", "iteration_3", "cutover")
_FILE_RE = re.compile(r"^Claude_BP_Effort_Output_(?P<client>.+?)_\d{8}_\d{6}(?:_Summary)?$", re.IGNORECASE)
_ISO_RE = re.compile(r"^[A-Z]{2}$")


class WorkbookFormatError(ValueError):
    """The file does not follow the effort-estimation template closely enough to be read."""


class DmTable(BaseModel):
    wave: str = "Wave 1"
    sequence: int = 1
    partition_axis: str = ""
    total_weeks: float = 0.0
    hypercare_weeks: float = 0.0
    countries: list[str] = Field(default_factory=list)
    rows: list[dict[str, Any]] = Field(default_factory=list)


class WorkbookSnapshot(BaseModel):
    source_file: str
    client_name: str = "Client"
    settings: dict[str, float] = Field(default_factory=dict)
    hypercare_days: float | None = None
    submodules: dict[str, list[str]] = Field(default_factory=dict)          # LOB -> sub-module names
    lines: list[EffortRow] = Field(default_factory=list)
    non_catalogue: list[NonCatalogueEffort] = Field(default_factory=list)
    tech_dev: list[TechDevRow] = Field(default_factory=list)
    dm_tables: list[DmTable] = Field(default_factory=list)
    basis: list[dict[str, Any]] = Field(default_factory=list)
    security: list[dict[str, Any]] = Field(default_factory=list)
    analytics: list[dict[str, Any]] = Field(default_factory=list)
    grids: list[WaveResourceGrid] = Field(default_factory=list)
    notes: list[str] = Field(default_factory=list)

    # ------------------------------------------------------------------ figures
    def effort(self, policy: Policy | None = None) -> EffortModel:
        policy = policy or get_policy()
        modules: dict[str, ModuleEffort] = {}
        for line in self.lines:
            module = modules.setdefault(line.lob, ModuleEffort(lob=line.lob, submodules=self.submodules.get(line.lob, [])))
            module.rows.append(line.model_copy())
        tables = WorkstreamTables(
            data_migration=[[dict(r) for r in t.rows] for t in self.dm_tables] or [[]],
            basis=[dict(r) for r in self.basis], security=[dict(r) for r in self.security],
            analytics=[{**r, **analytics_object_effort(r, policy)} for r in self.analytics])
        return EffortModel(rate_card_sheet=policy.effort.rate_card.default_sheet, modules=modules,
                           non_catalogue=[n.model_copy() for n in self.non_catalogue],
                           tech_dev=[t.model_copy() for t in self.tech_dev], workstreams=tables)

    def timeline(self) -> ResolvedTimeline:
        waves: list[ResolvedWave] = []
        for index, table in enumerate(self.dm_tables):
            if table.total_weeks > 0:
                waves.append(ResolvedWave(name=table.wave, sequence=table.sequence or index + 1,
                                          total_weeks=table.total_weeks, hypercare_weeks=table.hypercare_weeks,
                                          duration_source="workbook", countries=table.countries,
                                          partition_axis=table.partition_axis))
        known = {w.name for w in waves}
        for index, grid in enumerate(self.grids):   # waves only the resource grids show
            if grid.wave_name in known or not grid.months:
                continue
            pgls = sum(1 for band in grid.phase_bands if band == "PGLS")
            waves.append(ResolvedWave(name=grid.wave_name, sequence=index + 1, total_weeks=weeks_for(grid.months),
                                      hypercare_weeks=weeks_for(pgls), duration_source="workbook"))
        waves.sort(key=lambda w: w.sequence)
        hypercare = any(w.hypercare_weeks for w in waves)
        return ResolvedTimeline(waves=waves, sequencing="sequential" if len(waves) > 1 else "parallel",
                                hypercare_mode="inclusive" if hypercare else "none", source="rfp",
                                notes=["timeline read from the effort workbook"])

    def plan(self, policy: Policy | None = None) -> ResourcePlan:
        from bidcore.resourcing import price_plan

        policy = policy or get_policy()
        grids = [g.model_copy(deep=True) for g in self.grids]
        start = 0
        for grid in grids:              # waves stack on the programme calendar (sequential)
            grid.start_month = start
            start += grid.months
        plan = ResourcePlan(grids=grids, source="workbook")
        price_plan(plan, policy, daily_rate=self.settings.get("daily_rate_usd"))
        if self.hypercare_days is not None:
            plan.hypercare_effort_days = self.hypercare_days
        return plan

    def sizing(self, policy: Policy | None = None):
        from bidcore.sizing import SizingResult

        policy = policy or get_policy()
        effort = self.effort(policy)
        totals = workstream_totals(effort.workstreams, policy)
        plan = self.plan(policy)
        hypercare = self.hypercare_days if self.hypercare_days is not None else plan.hypercare_effort_days
        settings = {k: v for k, v in self.settings.items() if k != "hypercare_effort_days"}
        summary = compute_summary(effort.core_bp_effort + effort.non_catalogue_effort, totals["data_migration"],
                                  effort.tech_dev_effort, totals["security"], totals["basis"], totals["analytics"],
                                  hypercare, policy, settings)
        timeline = self.timeline()
        days = float(policy.commercials.working_days_per_month)
        names = [g.wave_name for g in plan.grids] or [w.name for w in timeline.waves] or ["Wave 1"]
        by_grid = {g.wave_name: round(g.grand_total_mm * days, 2) for g in plan.grids}
        wave_effort = WaveEffort(wave_names=names, lob=[{} for _ in names],
                                 pools=[{"Total": by_grid.get(n, 0.0)} for n in names],
                                 tech_dev=[0.0 for _ in names], category=[{} for _ in names],
                                 source={n: "workbook" for n in names})
        return SizingResult(effort=effort, timeline=timeline, wave_effort=wave_effort, weights=PassWeights(),
                            plan=plan, summary=summary, totals=totals, notes=list(self.notes))

    # ------------------------------------------------------------------ ledger
    def ledger(self, bid_id: str, policy: Policy | None = None) -> Ledger:
        policy = policy or get_policy()
        led = new_ledger(bid_id, self.client_name, policy_version=policy.version,
                         rate_card_sheet=policy.effort.rate_card.default_sheet)
        led.meta.source = "workbook"                     # Meta allows extra fields
        led.meta.source_workbook = self.source_file
        quote = lambda text: [Evidence(quote=str(text)[:300], file=self.source_file, page="workbook")]  # noqa: E731

        countries = sorted({line.country for line in self.lines if _ISO_RE.match(line.country or "")}
                           | {c for t in self.dm_tables for c in t.countries})
        allowed = set(policy.catalogue.allowed_countries)
        led.rfp_profile.data = RfpProfile(
            client_name=self.client_name, summary="Scope and effort as recorded in the reviewed effort workbook.",
            countries=[CountryScope(code=c) for c in countries if c in allowed])
        led.rfp_profile.state = "written"

        items: dict[tuple[str, str], ScopeItem] = {}
        for line in self.lines:
            key = (line.lob, line.scope_id)
            item = items.get(key)
            if item is None:
                item = items[key] = ScopeItem(row_id=f"si-{len(items) + 1}", scope_item_id=line.scope_id, lob=line.lob,
                                              business_area=line.business_area, description=line.description,
                                              mapping_basis="reviewer", written_by="workbook",
                                              submodule=line.submodule)
            if _ISO_RE.match(line.country or "") and line.country not in item.countries:
                item.countries.append(line.country)
        led.scope_items.rows = list(items.values())

        led.non_catalogue.rows = [NonCatalogueItem(row_id=f"nc-{i + 1}", name=n.name, label=n.label, description=n.description,
                                                   effort_days=n.effort_days, is_project_management=n.is_project_management,
                                                   written_by="workbook", evidence=quote(n.name))
                                  for i, n in enumerate(self.non_catalogue)]
        third = [t for t in self.tech_dev if t.source == "third_party"]
        led.integrations.rows = [Integration(row_id=f"int-{i + 1}", system=t.object_name, middleware=t.middleware,
                                             source_system=t.source_system, target_system=t.target_system,
                                             interface_count=max(t.no_of_objects, 1), complexity=_complexity(t.complexity),
                                             sap_modules=_sap_modules(t.source_system, t.target_system),
                                             effort_days=t.total, written_by="workbook", evidence=quote(t.object_name))
                                 for i, t in enumerate(third)]

        objects: dict[str, DataMigrationObject] = {}
        for table in self.dm_tables:
            for row in table.rows:
                key = str(row["label"]).strip().lower()
                if key in objects:
                    continue
                module, _, obj = str(row["label"]).partition(" - ")
                objects[key] = DataMigrationObject(
                    row_id=row["row_id"], object=(obj or module).strip(), module=module.strip() if obj else "",
                    rfp_label=str(row["label"]), status=row["status"], written_by="workbook",
                    **{k: float(row.get(k) or 0.0) for k in DM_KEYS})
        led.data_migration.rows = list(objects.values())
        led.basis.rows = [BasisActivity(row_id=r["row_id"], activity=r["activity"], status=r["status"],
                                        dev=_int(r.get("dev")), qa=_int(r.get("qa")), prd=_int(r.get("prd")),
                                        written_by="workbook") for r in self.basis]
        led.security.rows = [SecurityActivity(row_id=r["row_id"], activity=r["activity"], status=r["status"],
                                              complexity=_complexity(r.get("complexity")),
                                              effort_days=_int(r.get("effort_days")), written_by="workbook")
                             for r in self.security]
        led.analytics.rows = [AnalyticsObject(row_id=r["row_id"], object=r["object"], object_type=r["object_type"],
                                              build=r["build"], no_of_objects=max(int(r["no_of_objects"] or 1), 1),
                                              status=r["status"], complexity=_complexity(r.get("complexity")),
                                              written_by="workbook") for r in self.analytics]
        timeline = self.timeline()
        led.timeline.data = Timeline(
            waves=[Wave(name=w.name, sequence=w.sequence, total_weeks=w.total_weeks, hypercare_weeks=w.hypercare_weeks,
                        countries=[c for c in w.countries if c in allowed], duration_source="rfp",
                        partition_axis=w.partition_axis if w.partition_axis in ("entity", "country", "module", "site", "mixed") else "")
                   for w in timeline.waves],
            sequencing=timeline.sequencing, hypercare_required=timeline.hypercare_mode != "none",
            hypercare_mode=timeline.hypercare_mode, source="rfp", reason="read from the effort workbook")
        for section in ("scope_items", "non_catalogue", "integrations", "data_migration", "basis", "security",
                        "analytics", "timeline"):
            sec = getattr(led, section)
            filled = bool(sec.rows) if hasattr(sec, "rows") else sec.data is not None
            sec.state = "written" if filled else "empty"
            sec.none_reason = "" if filled else "not in the effort workbook"
            sec.written_by = ["workbook"]
        led.meta.status = "workbook_imported"
        return led


# ============================================================================ reading
def _num(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text or text.startswith("="):
        return None
    text = text.lstrip("~").split(" ")[0]
    try:
        return float(text)
    except ValueError:
        return None


def _int(value: Any) -> int | None:
    number = _num(value)
    return None if number is None else int(round(number))


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _complexity(value: Any) -> str:
    text = _text(value).lower()
    if text in ("high", "complex"):
        return "High"
    if text in ("low", "simple"):
        return "Low"
    return "Medium"


def _sap_modules(*sides: str) -> list[str]:
    found: list[str] = []
    for side in sides:
        for group in re.findall(r"SAP S/4HANA \(([^)]*(?:\([^)]*\))?[^)]*)\)", side or ""):
            for name in re.split(r",\s*", group):
                if name.strip() and name.strip() not in found:
                    found.append(name.strip())
    return found


class _Sheet:
    """One worksheet with formulas (values) and, when the file has them, Excel's cached results."""

    def __init__(self, ws, cached):
        self.ws, self.cached = ws, cached
        self.title = ws.title

    def raw(self, row: int, col: int) -> Any:
        return self.ws.cell(row=row, column=col).value

    def value(self, row: int, col: int) -> Any:
        value = self.raw(row, col)
        if isinstance(value, str) and value.startswith("=") and self.cached is not None:
            return self.cached.cell(row=row, column=col).value
        return value

    def number(self, row: int, col: int) -> float | None:
        return _num(self.value(row, col))

    def text(self, row: int, col: int) -> str:
        value = self.value(row, col)
        return "" if isinstance(value, str) and value.startswith("=") else _text(value)

    @property
    def max_row(self) -> int:
        return self.ws.max_row


def _is_total(label: str) -> bool:
    return label.lower().startswith("total")


def read_effort_workbook(path: Path, client_name: str = "") -> WorkbookSnapshot:
    """Every sheet of an effort workbook in the template layout -> WorkbookSnapshot."""
    path = Path(path)
    try:
        wb = load_workbook(path)                      # formulas
        cached = load_workbook(path, data_only=True)  # Excel's last results (None when saved by openpyxl)
    except Exception as exc:
        raise WorkbookFormatError(f"{path.name} is not a readable .xlsx workbook: {exc}") from exc
    snap = WorkbookSnapshot(source_file=path.name, client_name=client_name.strip() or _client_from(path, wb))
    try:
        for ws in wb.worksheets:
            sheet = _Sheet(ws, cached[ws.title] if ws.title in cached.sheetnames else None)
            a1, b1 = sheet.text(1, 1), sheet.text(1, 2)
            if a1 == "Line of Business":
                _read_lob(sheet, snap)
            elif a1 == "SAP Tool":
                _read_non_catalogue(sheet, snap)
            elif a1 == "Module" and b1.startswith("Object ID"):
                _read_tech_dev(sheet, snap)
            elif a1 == "Module" and b1 == "Business Area":
                _read_functional_scope(sheet, snap)
            elif a1.startswith("Summary") and "Effort" in a1 or ws.title == "Summary of Project Effort":
                _read_summary(sheet, snap)
            elif a1 == "Resource Estimation" or ws.title == "Project Timeline":
                _read_grids(sheet, snap)
            elif any(sheet.text(r, 1).startswith("Data Migration - Conversion Objects") for r in range(1, min(sheet.max_row, 12) + 1)):
                _read_data_migration(sheet, snap)
            elif sheet.text(2, 1) == "Activities":
                _read_basis(sheet, snap)
            elif sheet.text(2, 1) == "Activity" and sheet.text(2, 2).startswith("In Scope"):
                _read_security(sheet, snap)
            elif sheet.text(2, 1) == "Object" and sheet.text(2, 2) == "Object Type":
                _read_analytics(sheet, snap)
    finally:
        wb.close()
        cached.close()
    if not snap.lines and not snap.tech_dev and not snap.dm_tables and not snap.non_catalogue:
        raise WorkbookFormatError(
            f"{path.name} does not look like an effort workbook: no LOB sheet ('Line of Business' header), "
            "Tech Dev Scope, Data Migration Scope or Non Catalogue SAP Tools sheet was found")
    return snap


def _client_from(path: Path, wb) -> str:
    if "_bid" in wb.sheetnames:
        for row in wb["_bid"].iter_rows(min_row=1, max_col=2, values_only=True):
            if row and _text(row[0]) == "client" and _text(row[1]):
                return _text(row[1])
    m = _FILE_RE.match(path.stem)
    return m.group("client").replace("_", " ").strip() if m else "Client"


def _read_lob(s: _Sheet, snap: WorkbookSnapshot) -> None:
    for r in range(2, s.max_row + 1):
        scope_id, label = s.text(r, 4), s.text(r, 5)
        if label.upper() == "TOTAL" or (not scope_id and not s.text(r, 1)):
            continue
        if not scope_id:
            continue
        phases = {f: s.number(r, 7 + i) or 0.0 for i, f in enumerate(PHASES)}
        factor = s.number(r, 12)
        line = EffortRow(row_id=f"wb-{s.title}-{r}", item_row_id="", lob=s.text(r, 1) or s.title,
                         business_area=s.text(r, 2), country=s.text(r, 3).upper(), scope_id=scope_id,
                         description=label, complexity=s.text(r, 6), multiplication_factor=1.0 if factor is None else factor,
                         wave=s.text(r, 14), **phases)
        line.compute_total()
        snap.lines.append(line)


def _read_functional_scope(s: _Sheet, snap: WorkbookSnapshot) -> None:
    for r in range(2, s.max_row + 1):
        name = s.text(r, 1)
        m = re.match(r"^(.*?)\s*\[(.*)\]\s*$", name)
        if m:
            snap.submodules[m.group(1).strip()] = [x.strip() for x in re.split(r",\s*(?![^()]*\))", m.group(2)) if x.strip()]


def _read_non_catalogue(s: _Sheet, snap: WorkbookSnapshot) -> None:
    from bidcore.resourcing import is_project_management

    for r in range(2, s.max_row + 1):
        name = s.text(r, 1)
        if not name or name == "Non-Catalogue SAP Tools" or _is_total(name):
            continue
        effort = s.number(r, 4)
        if effort is None:
            continue
        description = s.text(r, 2)
        snap.non_catalogue.append(NonCatalogueEffort(
            row_id=f"nc-{len(snap.non_catalogue) + 1}", name=name, description=description, label=description,
            effort_days=effort, is_project_management=is_project_management(name)))


def _tech_source(name: str, object_type: str, source: str, target: str) -> str:
    kind = object_type.lower()
    if kind == "fiori" or "fiori" in name.lower():
        return "fiori"
    if kind == "interface":
        return "interface" if name.lower() in ("interface", "interfaces") or not (source or target) else "third_party"
    return "ricefw"


def _read_tech_dev(s: _Sheet, snap: WorkbookSnapshot) -> None:
    for r in range(2, s.max_row + 1):
        name = s.text(r, 2)
        if not name or _is_total(name) or _is_total(s.text(r, 1)):
            continue
        object_type, source, target = s.text(r, 3), s.text(r, 5), s.text(r, 6)
        snap.tech_dev.append(TechDevRow(
            row_id=f"td-{len(snap.tech_dev) + 1}", source=_tech_source(name, object_type, source, target),
            module=s.text(r, 1) or "All", object_name=name, object_type=object_type, middleware=s.text(r, 4),
            source_system=source, target_system=target, no_of_objects=_int(s.value(r, 7)) or 0,
            complexity=s.text(r, 8) or "Medium",
            **{f: round(s.number(r, 9 + i) or 0.0, 2)
               for i, f in enumerate(("development", "configuration", "unit_testing", "qa_testing"))}))


def _read_summary(s: _Sheet, snap: WorkbookSnapshot) -> None:
    labels = dict(SUMMARY_PARAMETERS)
    for r in range(1, s.max_row + 1):
        key = labels.get(s.text(r, 4))
        if key:
            value = s.number(r, 5)
            if value is not None and value >= 0:
                snap.settings[key] = value
        if s.text(r, 1) == HYPERCARE_LABEL:
            value = s.number(r, 2)
            if value is not None:
                snap.hypercare_days = value


def _read_grids(s: _Sheet, snap: WorkbookSnapshot) -> None:
    r = 1
    while r <= s.max_row:
        title = s.text(r, 1)
        if not title.startswith("Resource Plan - "):
            r += 1
            continue
        wave = title[len("Resource Plan - "):].strip() or f"Wave {len(snap.grids) + 1}"
        header = r + 1
        months = 0
        while s.text(header, 3 + months).upper().startswith("M") and s.text(header, 3 + months)[1:].isdigit():
            months += 1
        bands = [s.text(header + 1, 3 + i) for i in range(months)] if s.text(header + 1, 1) == "Phase" else [""] * months
        grid = WaveResourceGrid(wave_name=wave, months=months, phase_bands=bands)
        row = header + (2 if s.text(header + 1, 1) == "Phase" else 1)
        while row <= s.max_row:
            label = s.text(row, 1)
            if label == "Grand Total" or label.startswith("Resource Plan - ") or label.startswith("Reconciliation"):
                break
            if label and label not in ("Onsite Total", "Offshore Total"):
                location = OFFSHORE if s.text(row, 2).lower().startswith("off") else ONSITE
                cells = [GridCell(month_index=i + 1, phase=bands[i] if i < len(bands) else "", fte=s.number(row, 3 + i) or 0.0)
                         for i in range(months) if (s.number(row, 3 + i) or 0.0) > 0]
                grid.rows.append(GridRow(role_title=label, location=location, entity_kind="programme",
                                         contribution="additive", cells=cells))
            row += 1
        grid.recompute_totals()
        snap.grids.append(grid)
        r = row


def _read_data_migration(s: _Sheet, snap: WorkbookSnapshot) -> None:
    current = DmTable()
    in_table = False
    for r in range(1, s.max_row + 1):
        a = s.text(r, 1)
        if a == "Wave" and s.text(r, 2) == "Sequence":
            current = DmTable(wave=s.text(r + 1, 1) or f"Wave {len(snap.dm_tables) + 1}",
                              sequence=_int(s.value(r + 1, 2)) or len(snap.dm_tables) + 1,
                              partition_axis=s.text(r + 1, 3), total_weeks=s.number(r + 1, 4) or 0.0,
                              hypercare_weeks=s.number(r + 1, 5) or 0.0,
                              countries=[c.strip().upper() for c in re.split(r"[,;/ ]+", s.text(r + 1, 6))
                                         if _ISO_RE.match(c.strip().upper())])
            in_table = False
            continue
        if a == "Objects":
            in_table = True
            if current.rows:              # a second table without its own wave row
                current = DmTable(wave=f"Wave {len(snap.dm_tables) + 1}", sequence=len(snap.dm_tables) + 1)
            continue
        if not in_table:
            continue
        if _is_total(a):
            snap.dm_tables.append(current)
            in_table = False
            continue
        if not a or a.lower() in ("master data", "transactional data"):
            continue
        status = s.text(r, 2) or "In Scope"
        values = {k: s.number(r, 3 + i) for i, k in enumerate(DM_KEYS)}
        if all(v is None for v in values.values()):
            continue
        current.rows.append({"row_id": f"dm-{len(current.rows) + 1}", "object": a, "label": a,
                             "status": "Out of Scope" if status.lower().startswith("out") else "In Scope",
                             "sap_module": "", "category": "", "mm_compulsory": False,
                             **{k: (v or 0.0) for k, v in values.items()}})
    if in_table and current.rows:
        snap.dm_tables.append(current)


def _status(value: str) -> str:
    return "Out of Scope" if value.lower().startswith("out") else "In Scope"


def _read_basis(s: _Sheet, snap: WorkbookSnapshot) -> None:
    for r in range(3, s.max_row + 1):
        name = s.text(r, 1)
        if not name or _is_total(name):
            continue
        status = _status(s.text(r, 2))
        snap.basis.append({"row_id": f"bas-{len(snap.basis) + 1}", "activity": name, "status": status,
                           **{k: (s.number(r, 3 + i) if status == "In Scope" else None)
                              for i, k in enumerate(("dev", "qa", "prd"))}})


def _read_security(s: _Sheet, snap: WorkbookSnapshot) -> None:
    for r in range(3, s.max_row + 1):
        name = s.text(r, 1)
        if not name or _is_total(name):
            continue
        status = _status(s.text(r, 2))
        snap.security.append({"row_id": f"sec-{len(snap.security) + 1}", "activity": name, "sap_product": "",
                              "status": status, "complexity": _complexity(s.text(r, 3)),
                              "effort_days": s.number(r, 4) if status == "In Scope" else None})


def _read_analytics(s: _Sheet, snap: WorkbookSnapshot) -> None:
    for r in range(3, s.max_row + 1):
        name = s.text(r, 1)
        if not name or name.startswith("Total"):
            if name.startswith("Total"):
                break
            continue
        object_type = s.text(r, 2) if s.text(r, 2) in ("Model", "Report", "CDS View") else "Report"
        snap.analytics.append({"row_id": f"ana-{len(snap.analytics) + 1}", "object": name, "object_type": object_type,
                               "build": "Standard" if s.text(r, 3).lower().startswith("standard") else "Custom",
                               "no_of_objects": _int(s.value(r, 4)) or 1, "status": _status(s.text(r, 5)),
                               "complexity": _complexity(s.text(r, 6))})

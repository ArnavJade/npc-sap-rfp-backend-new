"""The effort workbook's static skeleton: sheets, title bands, headers, widths and prototype rows.

`scripts/build_templates.py` saves it as `assets/templates/template.xlsx`. The filler starts from that
file when it exists (so presales can restyle the workbook in Excel with no code change) and from
`build_template()` otherwise. Every table sheet follows one shape the filler relies on:

    row 1  title band            row 3  header row
    row 4  prototype data row    row 5  prototype total row

The filler copies the prototype rows' styles onto the generated rows and then clears them. The LOB
sheet is a single prototype ('_lob') copied once per LOB.
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from bidcore.policy import Policy, get_policy
from bidcore.render.workbook import styles as st
from bidcore.render.workbook.layout import HEADER_ROW, TableLayout, table_layouts

LOB_PROTOTYPE = "_lob"
PROTO_ROW, PROTO_TOTAL_ROW = HEADER_ROW + 1, HEADER_ROW + 2
TABLE_SHEETS = {"non_catalogue": "non_catalogue", "tech_dev": "tech_dev", "data_migration": "data_migration",
                "basis": "basis", "security": "security", "analytics": "analytics"}


def _title_band(ws: Worksheet, text: str, width: int, fill_=st.TITLE_FILL) -> None:
    ws.cell(row=1, column=1, value=text)
    for c in range(1, width + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font, cell.alignment = fill_, Font(bold=True, size=12), st.WRAP
    if width > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.row_dimensions[1].height = 22


def _prototype_rows(ws: Worksheet, layout: TableLayout) -> None:
    width = len(layout.columns)
    for c, column in enumerate(layout.columns, start=1):
        data = ws.cell(row=PROTO_ROW, column=c)
        data.border, data.alignment = st.BORDER, st.WRAP
        if column.number_format:
            data.number_format = column.number_format
        total = ws.cell(row=PROTO_TOTAL_ROW, column=c)
        total.border, total.font, total.fill = st.BORDER, st.BOLD, st.GREEN_FILL
        if column.number_format:
            total.number_format = column.number_format
    ws.cell(row=PROTO_ROW, column=layout.row_id_col).font = Font(color="808080", size=8)
    ws.column_dimensions[get_column_letter(layout.row_id_col)].hidden = True
    for c, column in enumerate(layout.columns, start=1):
        ws.column_dimensions[get_column_letter(c)].width = column.width
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)
    del width


def _table_sheet(wb: Workbook, title: str, band: str, layout: TableLayout) -> Worksheet:
    ws = wb.create_sheet(title)
    _title_band(ws, band, len(layout.columns))
    st.header_row(ws, HEADER_ROW, [c.header for c in layout.columns], st.NAVY_FILL, st.HEADER_FONT,
                  bordered=True, centered=True)
    ws.row_dimensions[HEADER_ROW].height = 45
    _prototype_rows(ws, layout)
    return ws


def build_template(policy: Policy | None = None) -> Workbook:
    policy = policy or get_policy()
    spec = policy.workbook["sheets"]
    layouts = table_layouts(policy.workbook)
    wb = Workbook()
    wb.remove(wb.active)

    summary = wb.create_sheet(spec["project_summary"]["title"])
    _title_band(summary, spec["project_summary"]["title"], 2, st.AMBER_FILL)
    st.header_row(summary, 2, ["Effort Category", "Effort (Person Days)"], st.NAVY_FILL, st.HEADER_FONT, bordered=True)
    for c, text in ((4, "Parameter"), (5, "Value")):
        cell = summary.cell(row=2, column=c, value=text)
        cell.fill, cell.font, cell.border = st.NAVY_FILL, st.HEADER_FONT, st.BORDER
    for col, width in (("A", 48), ("B", 22), ("C", 4), ("D", 32), ("E", 14)):
        summary.column_dimensions[col].width = width

    functional = wb.create_sheet(spec["functional_scope"]["title"])
    _title_band(functional, spec["functional_scope"]["title"], 4, st.HEADER_FILL)
    functional.cell(row=1, column=1).font = st.HEADER_FONT
    st.header_row(functional, HEADER_ROW, spec["functional_scope"]["headers"], st.HEADER_FILL, st.HEADER_FONT,
                  bordered=True, centered=True)
    for col, width in (("A", 34), ("B", 42), ("C", 14), ("D", 22)):
        functional.column_dimensions[col].width = width

    timeline = wb.create_sheet(spec["project_timeline"]["title"])
    _title_band(timeline, spec["project_timeline"]["title"], 8, st.GRID_FILL)
    timeline.column_dimensions["A"].width = 38
    timeline.column_dimensions["B"].width = 12
    timeline.column_dimensions["C"].width = 14

    lob = _table_sheet(wb, LOB_PROTOTYPE, "{lob}", layouts["catalogue"])
    lob.cell(row=1, column=1).fill = st.HEADER_FILL
    lob.cell(row=1, column=1).font = st.HEADER_FONT

    for key in ("non_catalogue", "tech_dev", "data_migration", "basis", "security", "analytics"):
        sheet = spec[key]
        _table_sheet(wb, sheet["title"], sheet.get("table_title", sheet["title"]), layouts[key])

    bid = wb.create_sheet(spec["bid"]["title"])
    bid.sheet_state = "hidden"
    bid.column_dimensions["A"].width = 18
    bid.column_dimensions["B"].width = 40
    return wb
